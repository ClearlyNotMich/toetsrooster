"""OWB Toetsrooster — desktop-app (Qt6). Start: dubbelklik op het programma.

Leest dezelfde Excel-werkmap als de Office Scripts en gebruikt dezelfde logica
(rooster_core.py). Het origineel wordt nooit overschreven: het resultaat wordt
als nieuw bestand opgeslagen (openpyxl zou de scriptknoppen verwijderen).
"""
from __future__ import annotations

import sys
from pathlib import Path

from PySide6.QtCore import QMarginsF, QTime, Qt
from PySide6.QtGui import QFont, QPageLayout, QPageSize, QPdfWriter, QTextDocument
from PySide6.QtWidgets import (QApplication, QButtonGroup, QCheckBox, QFileDialog, QFormLayout, QFrame, QGroupBox,
                               QHBoxLayout, QHeaderView, QLabel, QMainWindow, QMessageBox, QPushButton, QRadioButton,
                               QSpinBox, QTableWidget, QTableWidgetItem, QTabWidget, QTimeEdit, QVBoxLayout, QWidget)

import rooster_core as rc

STIJL = """
QWidget { color: #1f2d27; font-size: 10pt; }
QMainWindow, QWidget#midden, QStatusBar { background: #f5f7f6; }
QGroupBox { background: white; border: 1px solid #d6e0db; border-radius: 8px; margin-top: 14px; padding: 10px; font-weight: 600; }
QGroupBox::title { subcontrol-origin: margin; left: 10px; padding: 0 4px; color: #107c41; }
QPushButton { background: white; border: 1px solid #b9c9c0; border-radius: 6px; padding: 6px 14px; }
QPushButton:hover { background: #eaf4ee; }
QPushButton:disabled { color: #9aa8a0; background: #f0f2f1; }
QPushButton#hoofd { background: #107c41; color: white; border: none; font-weight: 600; padding: 9px 18px; }
QPushButton#hoofd:hover { background: #0b6a37; }
QPushButton#hoofd:disabled { background: #9cc8ae; }
QTabWidget::pane { border: 1px solid #d6e0db; border-radius: 6px; background: white; }
QTabBar::tab { padding: 7px 16px; background: #e7eeea; border-top-left-radius: 6px; border-top-right-radius: 6px; margin-right: 2px; }
QTabBar::tab:selected { background: white; color: #107c41; font-weight: 600; }
QTableWidget { background: white; gridline-color: #e3e9e6; alternate-background-color: #f7faf8; border: none; }
QHeaderView::section { background: #eef4f0; padding: 5px; border: none; border-bottom: 1px solid #d6e0db; font-weight: 600; }
QLabel#titel { font-size: 17pt; font-weight: 700; color: #107c41; }
QLabel#bestand { color: #4d5d55; }
QLabel#samenvatting { background: white; border: 1px solid #d6e0db; border-radius: 6px; padding: 8px; }
"""


def toon(kop: str, waarde) -> str:
    """Celwaarde netjes als tekst (datums en tijden zoals in Excel)."""
    k = rc.norm(kop)
    if isinstance(waarde, (int, float)) and not isinstance(waarde, bool) and waarde != "":
        if k in ("datum", "eerstedatum"):
            return rc.datum_tekst(int(waarde))
        if k in ("tijd", "eindtijd", "eerstetijd"):
            return rc.tijd_tekst(rc.naar_minuten(waarde))
    return rc.tekst(waarde)


def html_uitdraai(titel: str, matrix: list) -> str:
    rijen = "".join("<tr>" + "".join(f"<td>{_esc(c)}</td>" for c in rij) + "</tr>" for rij in matrix[1:])
    kop = "".join(f"<th>{_esc(c)}</th>" for c in matrix[0])
    return (f"<h2 style='color:#107c41'>{_esc(titel)}</h2><table width='100%' cellspacing='0' cellpadding='8' "
            f"border='1' style='border-collapse:collapse;border-color:#999'><tr style='background:#e7eeea'>{kop}</tr>{rijen}</table>")


def _esc(x) -> str:
    return rc.tekst(x).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def bewaar_pdf(html: str, pad: str) -> None:
    pdf = QPdfWriter(pad)
    pdf.setPageSize(QPageSize(QPageSize.A4))
    pdf.setPageOrientation(QPageLayout.Portrait)
    pdf.setResolution(300)
    pdf.setPageMargins(QMarginsF(15, 15, 15, 15), QPageLayout.Millimeter)
    doc = QTextDocument()
    doc.documentLayout().setPaintDevice(pdf)  # punten op PDF-resolutie, niet schermresolutie
    doc.setDefaultFont(QFont("Arial", 12))
    doc.setHtml(html)
    doc.setPageSize(pdf.pageLayout().paintRectPixels(pdf.resolution()).size().toSizeF())
    doc.print_(pdf)


def bewaar_werkmap(bron: str, doel: str, bladen: dict) -> None:
    from openpyxl import load_workbook
    wb = load_workbook(bron)
    rc.schrijf_bladen(wb, bladen)
    wb.save(doel)


class Venster(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("OWB Toetsrooster")
        self.resize(1180, 760)
        self.pad: str | None = None
        self.bladen: dict | None = None
        self.resultaat: dict | None = None

        midden = QWidget(objectName="midden")
        self.setCentralWidget(midden)
        hoofd = QVBoxLayout(midden)
        hoofd.setContentsMargins(16, 12, 16, 12)

        boven = QHBoxLayout()
        titel = QLabel("OWB Toetsrooster")
        titel.setObjectName("titel")
        boven.addWidget(titel)
        boven.addStretch()
        self.knop_open = QPushButton("📂  Werkmap openen…")
        self.knop_open.clicked.connect(self.kies_bestand)
        boven.addWidget(self.knop_open)
        hoofd.addLayout(boven)
        self.lbl_bestand = QLabel("Nog geen werkmap geopend. Open het Excel-bestand met de bladen Studenten, Slots, Beoordelaars en Simu.")
        self.lbl_bestand.setObjectName("bestand")
        self.lbl_bestand.setWordWrap(True)
        hoofd.addWidget(self.lbl_bestand)

        rij = QHBoxLayout()
        hoofd.addLayout(rij, 1)

        # ---- linkerkolom: keuzes
        links = QVBoxLayout()
        rij.addLayout(links, 0)
        g_modus = QGroupBox("Soort toets")
        lm = QVBoxLayout(g_modus)
        self.modus_groep = QButtonGroup(self)
        for i, (sleutel, m) in enumerate(rc.MODI.items()):
            rb = QRadioButton(f"{m['naam']}  ({m['slots']})")
            rb.setProperty("modus", sleutel)
            rb.setChecked(i == 0)
            self.modus_groep.addButton(rb)
            lm.addWidget(rb)
        links.addWidget(g_modus)

        g_inst = QGroupBox("Instellingen")
        f = QFormLayout(g_inst)
        self.jaar = QSpinBox(minimum=1, maximum=2)
        self.jaar.valueChanged.connect(self.toon_jaar)
        f.addRow("Studiejaar", self.jaar)
        self.duren = {}
        for label, standaard, uitleg in rc.INSTELLINGEN:
            if "(min)" in label:
                sb = QSpinBox(minimum=5, maximum=240, suffix=" min")
                sb.setValue(standaard)
                sb.setToolTip(uitleg)
                self.duren[label] = sb
                f.addRow(label.replace(" (min)", ""), sb)
        self.pauze = QCheckBox("Pauze in het rooster")
        self.pauze.setToolTip("Tijdens de pauze wordt niemand ingeroosterd; na de pauze geldt 'Casus na pauze'.")
        f.addRow(self.pauze)
        self.pauze_van = QTimeEdit(QTime(12, 0), displayFormat="HH:mm")
        self.pauze_tot = QTimeEdit(QTime(12, 30), displayFormat="HH:mm")
        self.pauze.toggled.connect(self.pauze_van.setEnabled)
        self.pauze.toggled.connect(self.pauze_tot.setEnabled)
        self.pauze.setChecked(False)
        self.pauze_van.setEnabled(False)
        self.pauze_tot.setEnabled(False)
        f.addRow("Pauze van", self.pauze_van)
        f.addRow("Pauze tot", self.pauze_tot)
        uitleg = QLabel("TTV-studenten krijgen de langere duur en komen bij voorkeur aan het eind van het blok vóór de pauze.")
        uitleg.setWordWrap(True)
        uitleg.setStyleSheet("color:#5b6b63; font-weight: normal;")
        f.addRow(uitleg)
        links.addWidget(g_inst)

        self.knop_maak = QPushButton("▶  Rooster maken")
        self.knop_maak.setObjectName("hoofd")
        self.knop_maak.setEnabled(False)
        self.knop_maak.clicked.connect(self.maak)
        links.addWidget(self.knop_maak)
        self.knop_opslaan = QPushButton("💾  Opslaan als Excel…")
        self.knop_opslaan.setToolTip("Slaat een kopie van de werkmap op met het rooster, de uitdraai en de meldingen.")
        self.knop_opslaan.clicked.connect(self.opslaan)
        self.knop_pdf = QPushButton("🖨  Uitdraai studenten (PDF)…")
        self.knop_pdf.setToolTip("Alleen naam, lokaal en tijd — geschikt om op te hangen of uit te delen.")
        self.knop_pdf.clicked.connect(self.pdf)
        for k in (self.knop_opslaan, self.knop_pdf):
            k.setEnabled(False)
            links.addWidget(k)
        links.addStretch()

        # ---- rechterkolom: resultaat
        rechts = QVBoxLayout()
        rij.addLayout(rechts, 1)
        self.samenvatting = QLabel("Open een werkmap en klik op ‘Rooster maken’.")
        self.samenvatting.setObjectName("samenvatting")
        self.samenvatting.setWordWrap(True)
        rechts.addWidget(self.samenvatting)
        self.tabs = QTabWidget()
        self.tabel_rooster = self.maak_tabel()
        self.tabel_uitdraai = self.maak_tabel()
        self.tabel_meldingen = self.maak_tabel()
        self.tabs.addTab(self.tabel_rooster, "Volledig rooster")
        self.tabs.addTab(self.tabel_uitdraai, "Uitdraai studenten")
        self.tabs.addTab(self.tabel_meldingen, "Meldingen")
        rechts.addWidget(self.tabs, 1)
        lijn = QFrame(frameShape=QFrame.HLine)
        lijn.setStyleSheet("color:#d6e0db")
        hoofd.addWidget(lijn)
        self.statusBar().showMessage("Klaar")
        self.toon_jaar()
        # Na een wijziging hoort het getoonde rooster niet meer bij de keuzes: eerst opnieuw maken.
        self.modus_groep.buttonToggled.connect(self.verouderd)
        for w in (self.jaar, *self.duren.values()):
            w.valueChanged.connect(self.verouderd)
        for w in (self.pauze_van, self.pauze_tot):
            w.timeChanged.connect(self.verouderd)
        self.pauze.toggled.connect(self.verouderd)

    def verouderd(self, *_):
        if self.resultaat:
            self.resultaat = None
            self.knop_opslaan.setEnabled(False)
            self.knop_pdf.setEnabled(False)
            self.samenvatting.setText("Instellingen gewijzigd — klik opnieuw op ‘Rooster maken’.")

    # ---------------------------------------------------------------- hulp

    @staticmethod
    def maak_tabel() -> QTableWidget:
        t = QTableWidget()
        t.setAlternatingRowColors(True)
        t.setEditTriggers(QTableWidget.NoEditTriggers)
        t.verticalHeader().setVisible(False)
        t.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeToContents)
        t.horizontalHeader().setStretchLastSection(True)
        return t

    def toon_jaar(self):
        j = self.jaar.value()
        for label, sb in self.duren.items():
            sb.setEnabled(f"jaar {j}" in label)

    def modus(self) -> str:
        return self.modus_groep.checkedButton().property("modus")

    def instellingen(self) -> dict:
        w = {"Jaar": self.jaar.value(), **{label: sb.value() for label, sb in self.duren.items()}}
        if self.pauze.isChecked():
            w["Pauze van"] = self.pauze_van.time().msecsSinceStartOfDay() / 86_400_000
            w["Pauze tot"] = self.pauze_tot.time().msecsSinceStartOfDay() / 86_400_000
        return w

    def vul(self, tabel: QTableWidget, matrix: list):
        kop, data = matrix[0], matrix[1:]
        tabel.clear()
        tabel.setColumnCount(len(kop))
        tabel.setRowCount(len(data))
        tabel.setHorizontalHeaderLabels([rc.tekst(k) for k in kop])
        for r, rij in enumerate(data):
            opvallend = any(rc.tekst(c) in ("NIET INGEROOSTERD", "Fout", "Waarschuwing") for c in rij)
            for c, waarde in enumerate(rij):
                item = QTableWidgetItem(toon(kop[c], waarde))
                if opvallend:
                    item.setBackground(Qt.GlobalColor.yellow if "Fout" not in rij else Qt.GlobalColor.red)
                tabel.setItem(r, c, item)

    # ---------------------------------------------------------------- acties

    def kies_bestand(self):
        pad, _ = QFileDialog.getOpenFileName(self, "Werkmap openen", "", "Excel-werkmap (*.xlsx *.xlsm)")
        if pad:
            self.open(pad)

    def open(self, pad: str) -> bool:
        try:
            bladen = rc.lees_werkmap(pad)
        except Exception as e:  # noqa: BLE001 - toon elke leesfout aan de gebruiker
            QMessageBox.critical(self, "Kan werkmap niet openen", f"{pad}\n\n{e}\n\nIs het bestand nog geopend in Excel?")
            return False
        self.pad, self.bladen = pad, bladen
        w = rc.instellingen_waarden(bladen.get("Instellingen"))
        jaar = rc.tekst(w["Jaar"])
        self.jaar.setValue(int(jaar) if jaar in ("1", "2") else 1)
        for label, sb in self.duren.items():
            if rc.tekst(w[label]).isdigit():
                sb.setValue(int(rc.tekst(w[label])))
        van, tot = rc.naar_minuten(w["Pauze van"]), rc.naar_minuten(w["Pauze tot"])
        self.pauze.setChecked(van is not None and tot is not None)
        if van is not None and tot is not None:
            self.pauze_van.setTime(QTime(van // 60, van % 60))
            self.pauze_tot.setTime(QTime(tot // 60, tot % 60))
        n = max(0, len(bladen.get("Studenten", [])) - 1)
        self.resultaat = None
        self.lbl_bestand.setText(f"<b>{Path(pad).name}</b> — {n} rijen in Studenten · bladen: {', '.join(bladen)}")
        self.knop_maak.setEnabled(True)
        for k in (self.knop_opslaan, self.knop_pdf):
            k.setEnabled(False)
        self.samenvatting.setText("Werkmap geladen. Controleer de instellingen en klik op ‘Rooster maken’.")
        self.statusBar().showMessage(f"Geopend: {pad}")
        return True

    def maak(self) -> dict:
        bladen = dict(self.bladen)
        bladen["Instellingen"] = rc.instellingen_matrix(self.instellingen())
        r = rc.maak_rooster(bladen, self.modus())
        r["bladen"]["Instellingen"] = bladen["Instellingen"]
        self.resultaat = r
        r["modus"] = self.modus()
        m = rc.MODI[r["modus"]]
        self.vul(self.tabel_meldingen, r["bladen"]["Meldingen"])
        if r["fouten"]:
            for t in (self.tabel_rooster, self.tabel_uitdraai):
                t.clear(), t.setRowCount(0), t.setColumnCount(0)
            self.tabs.setCurrentWidget(self.tabel_meldingen)
            self.samenvatting.setText(f"<b style='color:#b00020'>Er zijn {len(r['fouten'])} fout(en) in de werkmap.</b> "
                                      "Los ze op in Excel en open de werkmap opnieuw.")
        else:
            self.vul(self.tabel_rooster, r["bladen"][m["rooster"]])
            self.vul(self.tabel_uitdraai, r["bladen"][m["uitdraai"]])
            waarsch = sum(1 for x in r["meldingen"] if x[0] == "Waarschuwing")
            self.tabs.setCurrentWidget(self.tabel_rooster)
            kleur = "#b26a00" if waarsch else "#107c41"
            self.samenvatting.setText(f"<b style='color:{kleur}'>{r['meldingen'][0][1]}</b>"
                                      + (f"<br>{waarsch} waarschuwing(en) — zie tabblad Meldingen." if waarsch else ""))
        self.knop_opslaan.setEnabled(not r["fouten"])
        self.knop_pdf.setEnabled(not r["fouten"])
        return r

    def opslaan(self, doel: str | None = None) -> str | None:
        if not doel:
            voorstel = str(Path(self.pad).with_name(f"{Path(self.pad).stem} - {rc.MODI[self.resultaat['modus']]['rooster']}.xlsx"))
            doel, _ = QFileDialog.getSaveFileName(self, "Rooster opslaan als", voorstel, "Excel-werkmap (*.xlsx)")
            if not doel:
                return None
        if Path(doel).resolve() == Path(self.pad).resolve():
            QMessageBox.warning(self, "Kies een andere naam",
                                "Het origineel wordt niet overschreven (de scriptknoppen zouden verdwijnen). Kies een nieuwe bestandsnaam.")
            return None
        try:
            bewaar_werkmap(self.pad, doel, self.resultaat["bladen"])
        except Exception as e:  # noqa: BLE001
            QMessageBox.critical(self, "Opslaan mislukt", f"{doel}\n\n{e}")
            return None
        self.statusBar().showMessage(f"Opgeslagen: {doel}")
        return doel

    def pdf(self, doel: str | None = None) -> str | None:
        m = rc.MODI[self.resultaat["modus"]]
        if not doel:
            voorstel = str(Path(self.pad).with_name(f"{m['uitdraai']}.pdf"))
            doel, _ = QFileDialog.getSaveFileName(self, "Uitdraai opslaan als PDF", voorstel, "PDF (*.pdf)")
            if not doel:
                return None
        bewaar_pdf(html_uitdraai(f"Toetsrooster — {m['naam']}", self.resultaat["bladen"][m["uitdraai"]]), doel)
        self.statusBar().showMessage(f"PDF opgeslagen: {doel}")
        return doel


def controle(invoer: str, map_: str) -> int:
    """Zelftest zonder vensters (ook voor het gebouwde programma): --controle WERKMAP UITVOERMAP."""
    app = QApplication.instance() or QApplication(sys.argv)  # noqa: F841
    v = Venster()
    if not v.open(invoer):
        return 2
    code = 0
    for knop in v.modus_groep.buttons():
        knop.setChecked(True)
        r = v.maak()
        m = rc.MODI[v.modus()]
        doel = Path(map_) / f"controle-{m['rooster']}.xlsx"
        if r["fouten"] or not v.opslaan(str(doel)) or not v.pdf(str(doel.with_suffix(".pdf"))):
            code = 1
        if sys.stdout:  # None in het Windows-venster-programma
            print(f"{r['meldingen'][0][1] if not r['fouten'] else r['fouten']} -> {doel}")
    return code


def main() -> int:
    if len(sys.argv) == 4 and sys.argv[1] == "--controle":
        return controle(sys.argv[2], sys.argv[3])
    app = QApplication(sys.argv)
    app.setStyle("Fusion")
    app.setStyleSheet(STIJL)
    v = Venster()
    v.show()
    if len(sys.argv) == 2 and Path(sys.argv[1]).exists():
        v.open(sys.argv[1])
    return app.exec()


if __name__ == "__main__":
    sys.exit(main())
