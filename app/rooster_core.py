"""Roosterlogica voor OWB-toetsen (eerste kans en herkansing).

Dit is dezelfde logica als in de Office Scripts (`OWB * Roster Generator.txt`).
Wijzig je hier iets, wijzig het daar ook; tests/test_parity.py controleert dat
beide exact dezelfde uitvoer geven.

Invoer en uitvoer zijn "matrices": lijsten van rijen zoals Excel `getValues()`
ze geeft (datums als serienummers, tijden als dagfracties, lege cellen als "").
"""
from __future__ import annotations

import datetime as dt
import math
import re

KANS1 = "kans1"
HER = "her"
# Bladnamen komen uit de instellingen (sleutels hieronder).
MODI = {
    KANS1: {"naam": "Eerste kans", "dag": "blad_dag1", "slots": "blad_slots1", "rooster": "blad_rooster1",
            "uitdraai": "blad_uitdraai1"},
    HER: {"naam": "Herkansing", "dag": "blad_dag_her", "slots": "blad_slots_her", "rooster": "blad_rooster_her",
          "uitdraai": "blad_uitdraai_her"},
}
INSTELLINGEN_BLAD = "Instellingen"
INSTELLINGEN_KOP = ["Instelling", "Waarde", "Toelichting"]
ROOSTER_KOP = ["Datum", "Tijd", "Lokaal", "Casus", "Vaardigheid", "Beoordelaar", "simucliënt", "stnr",
               "Voornaam", "vv", "Achternaam", "klas", "bijzheden", "Eindtijd"]
UITDRAAI_KOP = ["Naam", "Lokaal", "Tijd"]
MELDINGEN_KOP = ["Soort", "Melding"]
EERSTE_KOP = ["EersteDatum", "EersteTijd", "EersteBeoordelaar", "EersteCasus"]
VRIJ = "Vrije plek"
NIET = "NIET INGEROOSTERD"
JA_NEE = ["Ja", "Nee"]
VOORKEUR = ["Voorkeur", "Verplicht", "Maakt niet uit"]
VOORKEUREN = ["Ochtend", "Middag", "Vroeg", "Laat"]
SIMU_WISSEL = ["Na elke pauze", "Na elke toets", "Nee"]
# Oude waarden in het blad Instellingen die nog geldig zijn.
ALIASSEN = {"simu_rotatie": {"ja": "Na elke pauze"}}
CASUS_NA = ["Nee", "Na pauze 1", "Na pauze 2", "Na pauze 3", "Na elke pauze afwisselen"]
GROEPEN = {
    "Toetsdag": "Uit deze tijden en het blad Toetsdag maakt het programma alle slots: om de toetsduur plus de "
                "wisseltijd begint een nieuwe student, behalve tijdens een pauze.",
    "Toetsduur": "Hoe lang een toets het lokaal, de beoordelaar en de simucliënt bezet houdt.",
    "Pauze": "Tijdens een pauze start of loopt geen gewone toets.",
    "TTV": "Studenten met toegestane tijdverlenging (kolom TTV = Ja, of 'ttv' in Opmerkingen).",
    "Voorkeur": "Studenten met een voorkeur (kolom Voorkeur, of ochtend/middag/vroeg/laat in Opmerkingen) "
                "krijgen eerder een plek, zo dicht mogelijk bij hun wens.",
    "Simucliënten": "Simucliënten uit het blad Simu (Datum, simucliënt, Vanaf, Tot).",
    "Beoordelaars": "Uitsluitingen uit de kolommen 'Niet inroosteren bij …' in het blad Studenten.",
    "Herkansing": "Wat er gebeurt als een herkanser opnieuw wordt ingeroosterd.",
    "Uitvoer": "Volgorde en inhoud van het rooster en de uitdraai voor studenten.",
    "Bladnamen": "Namen van de bladen in de werkmap. Alleen aanpassen als je bladen hebt hernoemd.",
}


def _inst(sleutel, groep, label, soort, standaard, uitleg, opties=(), min_=0, max_=0):
    return {"sleutel": sleutel, "groep": groep, "label": label, "soort": soort, "standaard": standaard,
            "uitleg": uitleg, "opties": list(opties), "min": min_, "max": max_}


# Eén bron voor de app, het blad Instellingen en de Office Scripts (tools/bouw_scripts.py zet deze lijst in de scripts).
INSTELLINGEN = [
    _inst("start", "Toetsdag", "Eerste toets om", "tijd", "09:40", "Begintijd van de eerste student."),
    _inst("eind", "Toetsdag", "Laatste toets klaar om", "tijd", "17:00", "Geen toets loopt later door dan deze tijd."),
    _inst("wissel", "Toetsdag", "Wisseltijd (min)", "getal", 5, "Tijd tussen twee studenten.", min_=0, max_=120),
    _inst("verdelen", "Toetsdag", "Lokalen vullen", "keuze", "Evenredig",
          "Evenredig: op elk tijdstip eerst het lokaal met de minste studenten, zodat elk lokaal (en elke beoordelaar) "
          "ongeveer evenveel krijgt, ook als er minder simu's dan lokalen zijn. Op volgorde: eerst de bovenste lokalen.",
          ["Evenredig", "Op volgorde"]),
    _inst("jaar", "Toetsduur", "Jaar", "getal", 1, "Studiejaar van deze toets; bepaalt welke duur hieronder geldt.", min_=1, max_=2),
    _inst("duur1", "Toetsduur", "Toetsduur jaar 1 (min)", "getal", 15, "Gewone toets in jaar 1.", min_=1, max_=480),
    _inst("ttv1", "Toetsduur", "TTV-toetsduur jaar 1 (min)", "getal", 20, "Toets met tijdverlenging in jaar 1.", min_=1, max_=480),
    _inst("duur2", "Toetsduur", "Toetsduur jaar 2 (min)", "getal", 15, "Gewone toets in jaar 2.", min_=1, max_=480),
    _inst("ttv2", "Toetsduur", "TTV-toetsduur jaar 2 (min)", "getal", 20, "Toets met tijdverlenging in jaar 2.", min_=1, max_=480),
    _inst("pauzes", "Pauze", "Pauzes", "pauzes", "11:15-11:40, 13:00-13:55, 15:15-15:40",
          "Van-tot, meerdere gescheiden door een komma. Leeg = geen pauze."),
    _inst("casus_na", "Pauze", "Casus na pauze", "keuze", "Na pauze 1",
          "Wanneer slots de casus uit de kolom 'Casus na pauze' krijgen (alleen als die kolom is ingevuld).", CASUS_NA),
    _inst("ttv_plek", "TTV", "TTV-plaatsing", "keuze", "Eind van blok vóór pauze",
          "Waar TTV-studenten bij voorkeur komen. Aan het eind van een blok valt de extra tijd in de pauze.",
          ["Eind van blok vóór pauze", "Zo vroeg mogelijk"]),
    _inst("ttv_eerst", "TTV", "TTV-studenten eerst inroosteren", "keuze", "Ja",
          "Geeft TTV-studenten als eerste een plek, zodat de beste plekken voor hen vrij zijn.", JA_NEE),
    _inst("middag", "Voorkeur", "Middag begint om", "tijd", "12:00", "Grens tussen voorkeur ochtend en middag."),
    _inst("simu_rotatie", "Simucliënten", "Simu's wisselen van lokaal", "keuze", "Na elke pauze",
          "Na elke pauze of na elke toets schuift elke simucliënt een lokaal door, zodat niemand de hele dag bij "
          "dezelfde toets zit. Nee: iedere simu blijft in hetzelfde lokaal.", SIMU_WISSEL),
    _inst("deelnaam", "Beoordelaars", "Namen deels vergelijken", "keuze", "Ja",
          "‘Ann’ sluit dan ook ‘Anneke’ uit. Vangt spelfouten op, maar kan te veel uitsluiten.", JA_NEE),
    _inst("her_beo", "Herkansing", "Andere beoordelaar dan 1e kans", "keuze", "Voorkeur",
          "Verplicht: lukt het niet, dan wordt de student niet ingeroosterd.", VOORKEUR),
    _inst("her_casus", "Herkansing", "Andere casus dan 1e kans", "keuze", "Voorkeur",
          "Verplicht: lukt het niet, dan wordt de student niet ingeroosterd.", VOORKEUR),
    _inst("vrije_plekken", "Uitvoer", "Vrije plekken tonen", "keuze", "Ja",
          "Lege maar bruikbare plekken als 'Vrije plek' in het volledige rooster (nooit op de uitdraai of in de PDF's).", JA_NEE),
    _inst("rooster_sort", "Uitvoer", "Rooster sorteren op", "keuze", "Lokaal, dan tijd",
          "Volgorde van het volledige rooster (altijd eerst per datum).", ["Lokaal, dan tijd", "Tijd, dan lokaal"]),
    _inst("uitdraai_sort", "Uitvoer", "Uitdraai sorteren op", "keuze", "Achternaam",
          "Volgorde van de uitdraai voor studenten.", ["Achternaam", "Tijd", "Lokaal, dan tijd"]),
    _inst("uitdraai_datum", "Uitvoer", "Datum op uitdraai", "keuze", "Ja",
          "Nee: alleen de begintijd (handig als alles op één dag is).", JA_NEE),
    _inst("blad_studenten", "Bladnamen", "Blad studenten", "tekst", "Studenten", ""),
    _inst("blad_dag1", "Bladnamen", "Blad toetsdag eerste kans", "tekst", "Toetsdag_Kans1", ""),
    _inst("blad_dag_her", "Bladnamen", "Blad toetsdag herkansing", "tekst", "Toetsdag_Her", ""),
    _inst("blad_simu", "Bladnamen", "Blad simucliënten", "tekst", "Simu", ""),
    _inst("blad_slots1", "Bladnamen", "Blad slots eerste kans (oud)", "tekst", "Slots_Kans1", ""),
    _inst("blad_slots_her", "Bladnamen", "Blad slots herkansing (oud)", "tekst", "Slots_Her", ""),
    _inst("blad_beo", "Bladnamen", "Blad beoordelaars (oud)", "tekst", "Beoordelaars", ""),
    _inst("blad_rooster1", "Bladnamen", "Blad rooster eerste kans", "tekst", "Rooster_Kans1", ""),
    _inst("blad_rooster_her", "Bladnamen", "Blad rooster herkansing", "tekst", "Rooster_Her", ""),
    _inst("blad_uitdraai1", "Bladnamen", "Blad uitdraai eerste kans", "tekst", "Uitdraai_Kans1", ""),
    _inst("blad_uitdraai_her", "Bladnamen", "Blad uitdraai herkansing", "tekst", "Uitdraai_Her", ""),
    _inst("blad_meldingen", "Bladnamen", "Blad meldingen", "tekst", "Meldingen", ""),
]
STANDAARD = {i["sleutel"]: i["standaard"] for i in INSTELLINGEN}


# ---------------------------------------------------------------- hulpjes

def norm(v) -> str:
    return re.sub(r"[^a-z0-9]", "", tekst(v).lower())


def tekst(v) -> str:
    if v is None or v == "":
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def waar(v) -> bool:
    if v is True or (not isinstance(v, bool) and isinstance(v, (int, float)) and v == 1):
        return True
    t = tekst(v).lower()
    return t.startswith("ja") or t in ("x", "waar", "true")


def serie(j: int, m: int, d: int) -> int | None:
    try:
        return (dt.date(j, m, d) - dt.date(1899, 12, 30)).days
    except ValueError:
        return None


MAAND = {"jan": 1, "feb": 2, "maa": 3, "mrt": 3, "mar": 3, "apr": 4, "mei": 5, "may": 5, "jun": 6, "jul": 7,
         "aug": 8, "sep": 9, "okt": 10, "oct": 10, "nov": 11, "dec": 12}
DATUM_VOORBEELD = "bijv. 14-04-2026, 14-4-26, 2026-04-14 of 14 april 2026"


def jaar4(j: str) -> int:
    return int(j) + 2000 if len(j) == 2 else int(j)


def naar_datum(v) -> int | None:
    """Excel-datum (getal) of tekst: 14-04-2026, 14/4/26, 14.04.2026, 2026-04-14, 2026/4/14,
    14 april 2026, di 14 apr 26, April 14, 2026 — eventueel met tijd erachter. Zonder jaartal: None."""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return math.floor(v) if v >= 1 else None
    t = re.sub(r"[\st]+\d{1,2}[:.]\d{2}([:.]\d{2})?$", "", tekst(v).lower()).strip()
    if re.fullmatch(r"\d{5}(\.\d+)?", t):
        return math.floor(float(t))
    m = re.fullmatch(r"(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", t)
    if m:
        return serie(int(m[1]), int(m[2]), int(m[3]))
    m = re.fullmatch(r"(\d{1,2})[-/.](\d{1,2})[-/.](\d{4}|\d{2})", t)
    if m:
        return serie(jaar4(m[3]), int(m[2]), int(m[1]))
    woorden = re.sub(r"[,.]", " ", t).split()
    m = re.fullmatch(r"(?:[a-z]+ )?(\d{1,2}) ([a-z]+) (\d{4}|\d{2})", " ".join(woorden))
    if m and m[2][:3] in MAAND:
        return serie(jaar4(m[3]), MAAND[m[2][:3]], int(m[1]))
    m = re.fullmatch(r"(?:[a-z]+ )?([a-z]+) (\d{1,2}) (\d{4}|\d{2})", " ".join(woorden))
    if m and m[1][:3] in MAAND:
        return serie(jaar4(m[3]), MAAND[m[1][:3]], int(m[2]))
    return None


def naar_minuten(v) -> int | None:
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return math.floor((v - math.floor(v)) * 1440 + 0.5)
    m = re.fullmatch(r"(\d{1,2})[:.](\d{2})", tekst(v))
    if m and int(m[1]) < 24 and int(m[2]) < 60:
        return int(m[1]) * 60 + int(m[2])
    return None


def datum_tekst(d: int) -> str:
    x = dt.date(1899, 12, 30) + dt.timedelta(days=d)
    return f"{x.day:02d}-{x.month:02d}-{x.year}"


def tijd_tekst(minuten: int) -> str:
    return f"{minuten // 60:02d}:{minuten % 60:02d}"


def kolom(kop: list, *namen: str) -> int:
    """Kolomnummer op kopnaam: eerst exact, anders de eerste kop die ermee begint ('Klas p2' -> klas)."""
    genormd = [norm(k) for k in kop]
    for n in namen:
        if n in genormd:
            return genormd.index(n)
    for n in namen:
        for i, g in enumerate(genormd):
            if g.startswith(n):
                return i
    return -1


def cel(rij: list, i: int):
    return rij[i] if 0 <= i < len(rij) else ""


def namen_lijst(v) -> list[str]:
    return [x.strip() for x in re.split(r"[,;/]", tekst(v)) if x.strip()]


def beperkt(naam: str, verboden: list[str], deels: bool) -> bool:
    n = naam.strip().lower()
    return any(n == v.lower() or (deels and (v.lower() in n or n in v.lower())) for v in verboden)


def leeg(rij: list) -> bool:
    return all(tekst(x) == "" for x in rij)


def heeft_woord(t: str, woord: str) -> bool:
    """Woord aan het begin van een woord in t ('laat' in 'laatste', niet in 'plaatsen')."""
    return re.search(r"(^|[^a-z])" + woord, t.lower()) is not None


def voorkeur_in(v) -> str:
    return next((w for w in VOORKEUREN if heeft_woord(tekst(v), w.lower())), "")


def lees_pauzes(v) -> list | None:
    """'11:15-11:40, 13:00-13:55' -> [[675, 700], [780, 835]]; None als het niet klopt."""
    lijst = []
    for deel in re.split(r"[,;]", tekst(v)):
        if not deel.strip():
            continue
        m = re.fullmatch(r"(\d{1,2}[:.]\d{2})\s*[-–]\s*(\d{1,2}[:.]\d{2})", deel.strip())
        a, b = (naar_minuten(m[1]), naar_minuten(m[2])) if m else (None, None)
        if a is None or b is None or a >= b:
            return None
        lijst.append([a, b])
    lijst.sort()
    if any(y[0] < x[1] for x, y in zip(lijst, lijst[1:])):
        return None
    return lijst


# ---------------------------------------------------------------- inlezen

def toelichting(i: dict) -> str:
    return i["uitleg"] + (f" Keuzes: {' / '.join(i['opties'])}." if i["opties"] else "")


def instellingen_waarden(matrix: list | None) -> dict:
    """Sleutel -> ruwe waarde uit blad Instellingen (standaardwaarde als de rij ontbreekt)."""
    waarden = dict(STANDAARD)
    labels = {norm(i["label"]): i["sleutel"] for i in INSTELLINGEN}
    for rij in (matrix or [])[1:]:
        if norm(cel(rij, 0)) in labels:
            waarden[labels[norm(cel(rij, 0))]] = cel(rij, 1)
    return waarden


def instellingen_matrix(waarden: dict | None = None) -> list:
    """Blad Instellingen met standaardwaarden of de gegeven sleutel -> waarde."""
    return [INSTELLINGEN_KOP] + [[i["label"], (waarden or {}).get(i["sleutel"], i["standaard"]), toelichting(i)]
                                 for i in INSTELLINGEN]


def instellingen_aangevuld(matrix: list | None) -> list | None:
    """Blad Instellingen met ontbrekende rijen toegevoegd, of None als er niets ontbreekt."""
    if not matrix:
        return instellingen_matrix()
    aanwezig = {norm(cel(rij, 0)) for rij in matrix[1:]}
    nieuw = [[i["label"], i["standaard"], toelichting(i)] for i in INSTELLINGEN if norm(i["label"]) not in aanwezig]
    return [list(r) for r in matrix] + nieuw if nieuw else None


def lees_instellingen(matrix: list | None, fouten: list) -> dict:
    w = instellingen_waarden(matrix)
    s = {}
    for i in INSTELLINGEN:
        k, v, label = i["sleutel"], w[i["sleutel"]], i["label"]
        if i["soort"] == "getal":
            s[k] = int(tekst(v)) if tekst(v).isdigit() else -1
            if not i["min"] <= s[k] <= i["max"]:
                fouten.append(f"Instellingen: '{label}' moet een geheel getal van {i['min']} t/m {i['max']} zijn.")
                s[k] = i["standaard"]
        elif i["soort"] == "tijd":
            s[k] = naar_minuten(v)
            if s[k] is None:
                fouten.append(f"Instellingen: '{label}' moet een tijd zijn (bijv. 09:40).")
                s[k] = naar_minuten(i["standaard"])
        elif i["soort"] == "pauzes":
            s[k] = lees_pauzes(v)
            if s[k] is None:
                fouten.append(f"Instellingen: '{label}' moet zoiets zijn als '11:15-11:40, 13:00-13:55' "
                              "(van vóór tot, niet overlappend) of leeg.")
                s[k] = []
        elif i["soort"] == "keuze":
            gevonden = [o for o in i["opties"] if norm(o) == norm(v)] or \
                [a for k_, a in ALIASSEN.get(k, {}).items() if k_ == norm(v)]
            s[k] = gevonden[0] if gevonden else i["standaard"]
            if not gevonden:
                fouten.append(f"Instellingen: '{label}' moet een van deze waarden zijn: {' / '.join(i['opties'])}.")
        else:
            s[k] = tekst(v) or i["standaard"]
            if not tekst(v):
                fouten.append(f"Instellingen: '{label}' mag niet leeg zijn.")
    s["duur"], s["ttv"] = s[f"duur{s['jaar']}"], s[f"ttv{s['jaar']}"]
    if s["start"] >= s["eind"]:
        fouten.append("Instellingen: 'Eerste toets om' moet vóór 'Laatste toets klaar om' liggen.")
    return s


def pauzes_tekst(pz: list) -> str:
    return ", ".join(f"{tijd_tekst(a)}-{tijd_tekst(b)}" for a, b in pz)


def lees_studenten(matrix: list | None, blad: str, fouten: list, meldingen: list) -> tuple[list, list]:
    if not matrix:
        fouten.append(f"Blad '{blad}' ontbreekt of is leeg.")
        return [], []
    kop = list(matrix[0])
    k = {
        "nr": kolom(kop, "studentnummer", "stnr"), "voornaam": kolom(kop, "voornaam", "roepnaam"),
        "tv": kolom(kop, "tussenvoegsel", "vv"), "achternaam": kolom(kop, "achternaam"), "klas": kolom(kop, "klas"),
        "niet_sb": kolom(kop, "nietinroosterenbijsbdocent"),
        "niet_vd": kolom(kop, "nietinroosterenbijdevaardighedendocent", "nietinroosterenbijvaardighedendocent"),
        "ttv": kolom(kop, "ttv"), "her": kolom(kop, "isherkanser", "herkanser"),
        "eb": kolom(kop, "eerstebeoordelaar"), "ec": kolom(kop, "eerstecasus"),
        "vk": kolom(kop, "voorkeur"), "opm": kolom(kop, "opmerkingen", "opmerking"),
    }
    for naam in ("voornaam", "achternaam"):
        if k[naam] < 0:
            fouten.append(f"Blad '{blad}': kolom '{naam.capitalize()}' ontbreekt.")
    if fouten:
        return [], kop
    lijst = []
    oud = []
    for i, rij in enumerate(matrix[1:], start=1):
        if tekst(cel(rij, k["voornaam"])) == "" and tekst(cel(rij, k["achternaam"])) == "":
            continue
        her = cel(rij, k["her"])
        if isinstance(her, (int, float)) and not isinstance(her, bool) and her > 1:
            oud.append(str(i + 1))
        st = {
            "i": i, "nr": tekst(cel(rij, k["nr"])), "voornaam": tekst(cel(rij, k["voornaam"])),
            "tv": tekst(cel(rij, k["tv"])), "achternaam": tekst(cel(rij, k["achternaam"])),
            "klas": tekst(cel(rij, k["klas"])),
            "verboden": namen_lijst(cel(rij, k["niet_sb"])) + namen_lijst(cel(rij, k["niet_vd"])),
            "ttv": waar(cel(rij, k["ttv"])), "her": waar(her),
            "eb": tekst(cel(rij, k["eb"])), "ec": tekst(cel(rij, k["ec"])),
            "voorkeur": voorkeur_in(cel(rij, k["vk"])), "opm": tekst(cel(rij, k["opm"])),
        }
        naam = f"{st['voornaam']} {st['achternaam']}"
        if tekst(cel(rij, k["vk"])) and not st["voorkeur"]:
            meldingen.append(("Waarschuwing", f"{naam}: voorkeur '{tekst(cel(rij, k['vk']))}' onbekend "
                              f"(gebruik {' / '.join(VOORKEUREN)}); genegeerd."))
        if st["opm"]:
            gelezen = (["TTV"] if heeft_woord(st["opm"], "ttv") else []) + \
                      ([f"voorkeur {voorkeur_in(st['opm']).lower()}"] if voorkeur_in(st["opm"]) else [])
            st["ttv"] = st["ttv"] or heeft_woord(st["opm"], "ttv")
            st["voorkeur"] = st["voorkeur"] or voorkeur_in(st["opm"])
            meldingen.append(("Info", f"Opmerking bij {naam}: '{st['opm']}' — " + (
                f"gelezen als {' en '.join(gelezen)}; controleer dit." if gelezen
                else "niet automatisch verwerkt; controleer dit zelf.")))
        lijst.append(st)
    if oud:
        meldingen.append(("Waarschuwing", f"{blad} rij {', '.join(oud)}: 'IsHerkanser' bevat een getal/datum "
                          "(restant van de oude scriptfout) en wordt gelezen als 'Nee'. Maak deze cellen leeg."))
    return lijst, kop


def lees_slots(matrix: list | None, blad: str, fouten: list) -> list:
    if not matrix:
        fouten.append(f"Blad '{blad}' ontbreekt of is leeg.")
        return []
    kop = list(matrix[0])
    kd, kt, kl = kolom(kop, "datum"), kolom(kop, "tijd"), kolom(kop, "lokaal")
    kc, kv, kn = kolom(kop, "casus"), kolom(kop, "vaardigheid"), kolom(kop, "casusnapauze")
    if min(kd, kt, kl) < 0:
        fouten.append(f"Blad '{blad}': kolommen 'Datum', 'Tijd' en 'Lokaal' zijn verplicht.")
        return []
    slots = []
    for i, rij in enumerate(matrix[1:], start=1):
        if leeg(rij):
            continue
        d, t = naar_datum(cel(rij, kd)), naar_minuten(cel(rij, kt))
        if d is None or t is None or tekst(cel(rij, kl)) == "":
            fouten.append(f"Blad '{blad}' rij {i + 1}: ongeldige of lege datum ('{tekst(cel(rij, kd))}'), tijd of lokaal "
                          f"(datum {DATUM_VOORBEELD}).")
            continue
        slots.append({"n": len(slots), "d": d, "t": t, "lokaal": tekst(cel(rij, kl)), "casus": tekst(cel(rij, kc)),
                      "casus_na": tekst(cel(rij, kn)), "vaardigheid": tekst(cel(rij, kv))})
    return slots


def lees_beschikbaar(matrix: list | None, blad: str, namen: tuple, fouten: list) -> dict:
    if not matrix:
        fouten.append(f"Blad '{blad}' ontbreekt of is leeg.")
        return {}
    kop = list(matrix[0])
    kd, kt, kn = kolom(kop, "datum"), kolom(kop, "tijd"), kolom(kop, *namen)
    kc, kl = kolom(kop, "casus"), kolom(kop, "lokaal")
    if min(kd, kt, kn) < 0:
        fouten.append(f"Blad '{blad}': kolommen 'Datum', 'Tijd' en de naamkolom zijn verplicht.")
        return {}
    beschikbaar: dict = {}
    for i, rij in enumerate(matrix[1:], start=1):
        naam = tekst(cel(rij, kn))
        if naam == "":
            continue
        d, t = naar_datum(cel(rij, kd)), naar_minuten(cel(rij, kt))
        if d is None or t is None:
            fouten.append(f"Blad '{blad}' rij {i + 1}: ongeldige datum ('{tekst(cel(rij, kd))}') of tijd "
                          f"(datum {DATUM_VOORBEELD}).")
            continue
        beschikbaar.setdefault(f"{d}_{t}", []).append(
            {"naam": naam, "casus": tekst(cel(rij, kc)), "lokaal": tekst(cel(rij, kl))})
    return beschikbaar


def tijdgrid(s: dict) -> list:
    """Begintijden van de dag: elke toetsduur + wisseltijd een student; na een pauze weer na de wisseltijd."""
    tijden, t = [], s["start"]
    while t + s["duur"] <= s["eind"]:
        botsing = [p for p in s["pauzes"] if t < p[1] and t + s["duur"] > p[0]]
        if botsing:
            t = botsing[0][1] + s["wissel"]
            continue
        tijden.append(t)
        t += s["duur"] + s["wissel"]
    return tijden


def heeft_data(matrix: list | None) -> bool:
    return bool(matrix) and any(not leeg(r) for r in matrix[1:])


def kenmerken(v) -> tuple:
    """'C1.36 (keuken)' -> ('C1.36', ['keuken']); 'K3 vrouw (keuken, rolstoel)' -> ('K3 vrouw', ['keuken', 'rolstoel'])."""
    t = tekst(v)
    tags = [x.strip() for m in re.findall(r"\(([^)]*)\)", t) for x in re.split(r"[,;/]", m) if x.strip()]
    return re.sub(r"\s*\([^)]*\)", "", t).strip(), tags


def past_bij(eisen: list, tags: list) -> bool:
    return all(norm(e) in [norm(x) for x in tags] for e in eisen)


def heel_woord(t: str, woord: str) -> bool:
    return re.search(r"(^|[^a-z0-9])" + re.escape(woord.lower().strip()) + r"($|[^a-z0-9])", t.lower()) is not None


def speelt(kan: list, casus: str) -> bool:
    """Kan een simu met 'Kan spelen' = kan (bijv. ['vrouw'] of ['P1', 'P2']) deze casus spelen? Leeg = alles."""
    return not kan or any(norm(k) == norm(casus) or heel_woord(casus, k) for k in kan)


def koppel_lokalen(lokalen: list, eisen: list) -> list:
    """Welk lokaal krijgt eenheid k (casus + beoordelaar)? Eenheden met eisen ('keuken') eerst, de meest beperkte
    voorop, bij voorkeur in hun eigen lokaal; daarna de rest op volgorde. Geeft per k een lokaalindex (of None)."""
    n = len(lokalen)
    uit = [None] * len(eisen)
    bezet: list = []
    opties = lambda k: [j for j in ([k] if k < n else []) + list(range(n)) if j not in bezet]  # noqa: E731
    met_eis = sorted([k for k in range(len(eisen)) if eisen[k]],
                     key=lambda k: (sum(1 for j in range(n) if past_bij(eisen[k], lokalen[j][1])), k))
    for k in met_eis:
        j = next((j for j in opties(k) if past_bij(eisen[k], lokalen[j][1])), None)
        if j is not None:
            uit[k] = j
            bezet.append(j)
    for k in range(len(eisen)):
        if uit[k] is None and opties(k):
            uit[k] = opties(k)[0]
            bezet.append(uit[k])
    return uit


def meld_koppeling(lokalen: list, casussen: list, eisen: list, uit: list, d, meldingen: list) -> None:
    for k, j in enumerate(uit):
        if j is None:
            continue
        if eisen[k] and not past_bij(eisen[k], lokalen[j][1]):
            meldingen.append(("Waarschuwing", f"Casus {casussen[k]} vraagt {', '.join(eisen[k])}, maar op {datum_tekst(d)} "
                              f"is daar geen vrij lokaal voor; staat nu in {lokalen[j][0]}."))
        elif j != k and k < len(lokalen):
            meldingen.append(("Info", f"Casus {casussen[k]} staat in lokaal {lokalen[j][0]} "
                              f"({', '.join(eisen[k]) or 'ruil'}) in plaats van {lokalen[k][0]}."))


def lees_toetsdag(matrix: list, blad: str, s: dict, fouten: list, meldingen: list) -> tuple[list, dict]:
    """Blad Toetsdag (één rij per lokaal per dag) -> slots en beschikbare beoordelaars per slot.
    Beoordelaar, Beoordelaar 2, 3, …: per slot de eerste die dan kan. Tijden in de cel ('Fien tot 13.00',
    'Dorus 10.00-12.00') of met 'Wissel om' tussen beoordelaar 1 en 2.
    Kenmerken tussen haakjes: 'C1.36 (keuken)' en casus 'K3 (keuken)' -> die casus (met beoordelaar) gaat naar dat lokaal."""
    kop = list(matrix[0])
    kd, kl, kb = kolom(kop, "datum"), kolom(kop, "lokaal"), kolom(kop, "beoordelaar")
    kc, kn, kv = kolom(kop, "casus"), kolom(kop, "casusnapauze"), kolom(kop, "vaardigheid")
    kw = kolom(kop, "wisselom")
    kbs = [kb] + [i for i, k in enumerate(kop) if re.fullmatch(r"beoordelaar\d+", norm(k))]
    if min(kd, kl, kb) < 0:
        fouten.append(f"Blad '{blad}': kolommen 'Datum', 'Lokaal' en 'Beoordelaar' zijn verplicht.")
        return [], {}
    eenheden, gezien = [], []
    for i, rij in enumerate(matrix[1:], start=1):
        if leeg(rij):
            continue
        d, lokaal = naar_datum(cel(rij, kd)), kenmerken(cel(rij, kl))
        wie = [lees_inzet(tekst(cel(rij, k)), s["middag"]) for k in kbs if tekst(cel(rij, k))]
        w = naar_minuten(cel(rij, kw)) if tekst(cel(rij, kw)) else None
        if d is None or not lokaal[0] or not wie:
            fouten.append(f"Blad '{blad}' rij {i + 1}: vul een geldige datum ({DATUM_VOORBEELD}; nu "
                          f"'{tekst(cel(rij, kd))}'), een lokaal en een beoordelaar in.")
            continue
        if tekst(cel(rij, kw)) and (w is None or len(wie) < 2):
            fouten.append(f"Blad '{blad}' rij {i + 1}: 'Wissel om' moet een tijd zijn en vraagt ook een 'Beoordelaar 2'.")
            continue
        if w is not None:  # wissel: beoordelaar 1 tot w, beoordelaar 2 vanaf w (tenzij in de cel al een tijd staat)
            wie[0]["tot"] = wie[0]["tot"] if wie[0]["tot"] is not None else w
            wie[1]["van"] = wie[1]["van"] if wie[1]["van"] is not None else w
        elif len(wie) > 1 and all(p["van"] is None and p["tot"] is None for p in wie):
            fouten.append(f"Blad '{blad}' rij {i + 1}: bij meer beoordelaars in één lokaal: vul 'Wissel om' in, of zet "
                          "tijden in de cel (bijv. 'Fien tot 13.00' en 'Hans vanaf 13.30').")
            continue
        if f"{d}_{norm(lokaal[0])}" in gezien:
            fouten.append(f"Blad '{blad}' rij {i + 1}: lokaal {lokaal[0]} staat al eerder op {datum_tekst(d)}.")
            continue
        gezien.append(f"{d}_{norm(lokaal[0])}")
        eenheden.append({"d": d, "lokaal": lokaal, "casus": kenmerken(cel(rij, kc)), "casus_na": kenmerken(cel(rij, kn)),
                         "vaardigheid": tekst(cel(rij, kv)), "wie": wie})
    slots, beo = [], {}
    tijden = tijdgrid(s)
    for d in list(dict.fromkeys(e["d"] for e in eenheden)):
        rijen = [e for e in eenheden if e["d"] == d]
        lokalen = [e["lokaal"] for e in rijen]
        eisen = [e["casus"][1] for e in rijen]
        plek = koppel_lokalen(lokalen, eisen)
        meld_koppeling(lokalen, [e["casus"][0] for e in rijen], eisen, plek, d, meldingen)
        for k, e in enumerate(rijen):
            kamer = lokalen[plek[k]]
            if e["casus_na"][1] and not past_bij(e["casus_na"][1], kamer[1]):
                meldingen.append(("Waarschuwing", f"Casus na pauze {e['casus_na'][0]} vraagt {', '.join(e['casus_na'][1])}, "
                                  f"maar lokaal {kamer[0]} heeft dat niet."))
            e["kamer"] = kamer[0]
            for t in tijden:
                slots.append({"n": len(slots), "d": d, "t": t, "lokaal": kamer[0], "casus": e["casus"][0],
                              "casus_na": e["casus_na"][0], "vaardigheid": e["vaardigheid"]})
                past_ = [p for p in e["wie"] if (p["van"] is None or p["van"] <= t) and (p["tot"] is None or t + s["duur"] <= p["tot"])]
                if past_:
                    beo[f"{d}_{t}_{norm(kamer[0])}"] = [{"naam": past_[0]["naam"], "casus": "", "lokaal": ""}]
        meld_dubbel_ingezet(rijen, s, d, meldingen)
    return slots, beo


def meld_dubbel_ingezet(rijen: list, s: dict, d: int, meldingen: list) -> None:
    """Waarschuwing als een beoordelaar op dezelfde tijd in twee lokalen staat (bijv. bij ruilen)."""
    inzet = []
    for e in rijen:
        for p in e["wie"]:
            a = p["van"] if p["van"] is not None else s["start"]
            b = p["tot"] if p["tot"] is not None else s["eind"]
            for naam, k, x, y in inzet:
                if naam == norm(p["naam"]) and k != e["kamer"] and a < y and x < b:
                    meldingen.append(("Waarschuwing", f"{p['naam']} staat op {datum_tekst(d)} van {tijd_tekst(max(a, x))} tot "
                                      f"{tijd_tekst(min(b, y))} in twee lokalen ({k} en {e['kamer']}); hij of zij kan maar "
                                      "in één tegelijk beoordelen."))
            inzet.append((norm(p["naam"]), e["kamer"], a, b))


def lees_simu_dag(matrix: list, blad: str, fouten: list, meldingen: list) -> list:
    """Blad Simu zonder kolom Tijd: één rij per simucliënt per dag (Datum leeg = elke dag), met Vanaf/Tot en
    optioneel 'Vast lokaal': leeg = wisselt volgens de instelling, Ja = blijft in één lokaal, of een lokaalnaam."""
    kop = list(matrix[0])
    kn, kd = kolom(kop, "simuclint", "simu", "simulant", "naam"), kolom(kop, "datum")
    kv, kt = kolom(kop, "vanaf", "van"), kolom(kop, "tot")
    kl, kk = kolom(kop, "vastlokaal"), kolom(kop, "kanspelen")
    if kn < 0:
        fouten.append(f"Blad '{blad}': kolom 'simucliënt' ontbreekt.")
        return []
    lijst = []
    for i, rij in enumerate(matrix[1:], start=1):
        naam = tekst(cel(rij, kn))
        if not naam:
            continue
        d = naar_datum(cel(rij, kd)) if tekst(cel(rij, kd)) else None
        van = naar_minuten(cel(rij, kv)) if tekst(cel(rij, kv)) else None
        tot = naar_minuten(cel(rij, kt)) if tekst(cel(rij, kt)) else None
        if (tekst(cel(rij, kd)) and d is None) or (tekst(cel(rij, kv)) and van is None) or (tekst(cel(rij, kt)) and tot is None):
            fouten.append(f"Blad '{blad}' rij {i + 1}: ongeldige datum ('{tekst(cel(rij, kd))}'), Vanaf of Tot "
                          f"(datum {DATUM_VOORBEELD}).")
            continue
        dubbel = [x for x in lijst if norm(x["naam"]) == norm(naam) and x["d"] == d]
        if dubbel:
            meldingen.append(("Waarschuwing", f"Blad '{blad}' rij {i + 1}: {naam} staat al eerder op "
                              f"{datum_tekst(d) if d is not None else 'elke dag'}; telt maar één keer."))
            continue
        vast = tekst(cel(rij, kl))
        lijst.append({"naam": naam, "d": d, "van": van, "tot": tot, "kan": namen_lijst(cel(rij, kk)),
                      "vast": "" if norm(vast) in ("", "nee", "n", "wisselt") else "ja" if waar(vast) else vast})
    return lijst


def simu_rotatie(slots: list, simus: list, s: dict, meldingen: list) -> dict:
    """Simu per slot. Lokaal k van die dag krijgt bij voorkeur simu (k + stap) mod n, uit iedereen die in dat blok
    werkt; stap = aantal pauzes ervoor ('Na elke pauze') of het hoeveelste tijdstip van de dag ('Na elke toets').
    Per slot telt alleen wie dan volgens Vanaf/Tot beschikbaar is."""
    blok = lambda t: sum(1 for p in s["pauzes"] if p[1] <= t)  # noqa: E731
    kan = lambda x, a, e: (x["van"] is None or x["van"] <= a) and (x["tot"] is None or x["tot"] >= e)  # noqa: E731
    kamers: dict = {}
    bereik: dict = {}
    for slot in slots:
        lijst = kamers.setdefault(slot["d"], [])
        if norm(slot["lokaal"]) not in lijst:
            lijst.append(norm(slot["lokaal"]))
        sleutel = f"{slot['d']}_{blok(slot['t'])}"
        a, e, ks = bereik.get(sleutel, [slot["t"], slot["t"] + s["duur"], []])
        if norm(slot["lokaal"]) not in ks:
            ks.append(norm(slot["lokaal"]))
        bereik[sleutel] = [min(a, slot["t"]), max(e, slot["t"] + s["duur"]), ks]
    tijden = {d: sorted({x["t"] for x in slots if x["d"] == d}) for d in kamers}
    stap = lambda d, t: (blok(t) if s["simu_rotatie"] == "Na elke pauze"  # noqa: E731
                         else tijden[d].index(t) if s["simu_rotatie"] == "Na elke toets" else 0)
    for d in kamers:
        if not any(x["d"] is None or x["d"] == d for x in simus):
            andere = sorted({datum_tekst(x["d"]) for x in simus if x["d"] is not None})
            meldingen.append(("Waarschuwing", f"Geen simucliënten voor {datum_tekst(d)} in het blad Simu" +
                              (f" (er staan er wel voor {', '.join(andere)})." if andere else ".")))
    # Vaste simu's: eerst wie een lokaal heeft opgegeven, dan 'Ja' in het eerste lokaal dat nog vrij is.
    thuis: dict = {}
    for d, lokalen in kamers.items():
        thuis[d] = {}
        die_dag = []
        for x in simus:
            if (x["d"] is None or x["d"] == d) and x["vast"] and x["naam"] not in [y["naam"] for y in die_dag]:
                die_dag.append(x)
        for x in [y for y in die_dag if y["vast"] != "ja"]:
            if norm(x["vast"]) not in lokalen:
                meldingen.append(("Waarschuwing", f"{x['naam']}: vast lokaal '{x['vast']}' bestaat niet op "
                                  f"{datum_tekst(d)}; deze simu wisselt nu gewoon mee."))
            elif norm(x["vast"]) in thuis[d]:
                meldingen.append(("Waarschuwing", f"{x['naam']}: lokaal {x['vast']} is al vast voor "
                                  f"{thuis[d][norm(x['vast'])]['naam']}; deze simu wisselt nu gewoon mee."))
            else:
                thuis[d][norm(x["vast"])] = x
        for x in [y for y in die_dag if y["vast"] == "ja"]:
            vrij_ = [k for k in lokalen if k not in thuis[d]]
            if vrij_:
                thuis[d][vrij_[0]] = x
    vaste_namen = {d: [x["naam"] for x in h.values()] for d, h in thuis.items()}
    gekozen: dict = {}
    for sleutel, (a, e, ks) in bereik.items():
        d, b = (int(x) for x in sleutel.split("_"))
        geschikt = []
        for x in simus:
            werkt = (x["van"] is None or x["van"] < e) and (x["tot"] is None or x["tot"] > a)
            if (x["d"] is None or x["d"] == d) and werkt and x["naam"] not in [y["naam"] for y in geschikt]:
                geschikt.append(x)
        if geschikt and len(geschikt) < len(ks):
            meldingen.append(("Waarschuwing", f"{datum_tekst(d)} {tijd_tekst(a)}-{tijd_tekst(e)}: {len(geschikt)} "
                              f"simucliënt(en) beschikbaar voor {len(ks)} lokalen; niet elk lokaal kan dan draaien."))
        gekozen[sleutel] = geschikt
    uit = {}
    zonder = 0
    for slot in slots:
        b = blok(slot["t"])
        g = gekozen[f"{slot['d']}_{b}"]
        lokaal = norm(slot["lokaal"])
        h = thuis[slot["d"]].get(lokaal)
        eigen = [h] if h and h in g else []
        wissel = [x for x in g if x["naam"] not in vaste_namen[slot["d"]]]
        kamers_rot = [k_ for k_ in kamers[slot["d"]] if k_ not in thuis[slot["d"]]]
        plek = kamers_rot.index(lokaal) if lokaal in kamers_rot else kamers[slot["d"]].index(lokaal)
        k = (plek + stap(slot["d"], slot["t"])) % len(wissel) if wissel else 0
        lijst = [{"naam": x["naam"], "casus": "", "lokaal": ""} for x in eigen + wissel[k:] + wissel[:k]
                 if kan(x, slot["t"], slot["t"] + s["duur"]) and speelt(x["kan"], slot["casus"])]
        if lijst:
            uit[f"{slot['d']}_{slot['t']}_{norm(slot['lokaal'])}"] = lijst
        else:
            zonder += 1
    if zonder and any(gekozen.values()):
        meldingen.append(("Info", f"{zonder} slot(s) zonder beschikbare simucliënt (vóór 'Vanaf', na 'Tot', of niemand die "
                          "die casus kan spelen); daar wordt niemand ingeroosterd."))
    gemeld = []
    for slot in slots:  # casus die niemand van die dag kan spelen, en vaste simu's die hun casus niet kunnen spelen
        die_dag = [x for x in simus if x["d"] is None or x["d"] == slot["d"]]
        sleutel = f"{slot['d']}_{slot['casus']}"
        if die_dag and sleutel not in gemeld and not any(speelt(x["kan"], slot["casus"]) for x in die_dag):
            gemeld.append(sleutel)
            meldingen.append(("Waarschuwing", f"Casus {slot['casus']} op {datum_tekst(slot['d'])}: geen simucliënt die "
                              "deze casus kan spelen (kolom 'Kan spelen' in het blad Simu)."))
        h = thuis[slot["d"]].get(norm(slot["lokaal"]))
        sleutel = f"{slot['d']}_{h['naam'] if h else ''}_{slot['casus']}"
        if h and not speelt(h["kan"], slot["casus"]) and sleutel not in gemeld:
            gemeld.append(sleutel)
            meldingen.append(("Waarschuwing", f"{h['naam']} staat vast in {slot['lokaal']}, maar kan casus {slot['casus']} "
                              "niet spelen; daar spelen dan de anderen."))
    return uit


# ---------------------------------------------------------------- conceptblad (lijstvorm)
# Eén blad met in kolom A een kopje (datum, inzet, tijden, pauzes, lokalen, casussen, simu's, studenten) en in
# kolom B de waarden eronder, zoals de roostermaker het opschrijft. Wordt gebruikt als het blad Toetsdag leeg is.

SECTIES = [("casussennapauze", "casus_na"), ("casusnapauze", "casus_na"), ("casus", "casus"), ("vaardig", "vaardigheid"),
           ("datum", "datum"), ("inzet", "inzet"), ("beoordelaar", "inzet"), ("tijd", "tijden"), ("pauze", "pauzes"),
           ("lokal", "lokalen"), ("simu", "simus"), ("student", "studenten")]


def sectie(label: str) -> str:
    n = norm(label)
    return next((naam for begin, naam in SECTIES if n.startswith(begin)), "")


def vind_concept(bladen: dict) -> str:
    """Naam van het eerste blad in conceptvorm (kopjes 'lokalen' en 'tijden' of 'inzet' in kolom A), anders ''."""
    for naam, rijen in bladen.items():
        kopjes = {sectie(cel(r, 0)) for r in rijen or [] if tekst(cel(r, 0))}
        if "lokalen" in kopjes and ("tijden" in kopjes or "inzet" in kopjes):
            return naam
    return ""


def lees_inzet(t: str, middag: int) -> dict:
    """'Fien ochtend tot 13.00' -> naam Fien, tot 13:00; 'Dorus 10.00-12.00' of 'Dorus van 10.00 tot 12.00' -> van/tot;
    'Jet (nog niet eerder gedaan)' -> opmerking. 'Ida van Berkel' blijft gewoon een naam."""
    laag = t.lower()
    tijd = r"(\d{1,2}(?:[.:]\d{2})?)"
    uur = lambda x: naar_minuten(x if re.search(r"[.:]", x) else x + ":00")  # noqa: E731
    bereik = re.search(r"(\d{1,2}[.:]\d{2})\s*[-–]\s*(\d{1,2}[.:]\d{2})", laag)
    tot = re.search(r"\btot\s+" + tijd, laag)
    van = re.search(r"\b(?:vanaf|va|van)\s+" + tijd, laag)
    opm = re.search(r"\(([^)]*)\)", t)
    naam = re.split(r"\s*(?:\(|,|\b(?:ochtend|middag|tot|vanaf|va)\b|\bvan\s+\d|\d{1,2}[.:]\d{2})", t, maxsplit=1,
                    flags=re.IGNORECASE)[0].strip()
    if bereik:
        return {"naam": naam or t.strip(), "tekst": t.strip(), "opm": opm[1].strip() if opm else "",
                "van": uur(bereik[1]), "tot": uur(bereik[2])}
    return {"naam": naam or t.strip(), "tekst": t.strip(), "opm": opm[1].strip() if opm else "",
            "tot": uur(tot[1]) if tot else (middag if heeft_woord(laag, "ochtend") else None),
            "van": uur(van[1]) if van else (middag if heeft_woord(laag, "middag") else None)}


def wie_tekst(p: dict) -> str:
    if p["van"] is not None and p["tot"] is not None:
        return f"{p['naam']} ({tijd_tekst(p['van'])}-{tijd_tekst(p['tot'])})"
    return p["naam"] + (f" (tot {tijd_tekst(p['tot'])})" if p["tot"] is not None else "") + \
        (f" (vanaf {tijd_tekst(p['van'])})" if p["van"] is not None else "")


def lees_concept(matrix: list, blad: str, s: dict, fouten: list, meldingen: list) -> dict:
    """Conceptblad -> slots, beoordelaars per slot, aantal simu's nodig, eventueel studenten.
    Tijden, pauzes en toetsduur uit het blad gaan voor de instellingen (s wordt aangepast)."""
    waarden: dict = {}
    huidig = ""
    for rij in matrix:
        if tekst(cel(rij, 0)):
            huidig = sectie(cel(rij, 0))
        if huidig and tekst(cel(rij, 1)):
            waarden.setdefault(huidig, []).append(cel(rij, 1))
    uit = {"slots": [], "beo": {}, "nodig": 0, "studenten": None, "d": None}
    datums = waarden.get("datum", [])
    d = naar_datum(datums[0]) if datums else None
    if d is None:
        fouten.append(f"Blad '{blad}': vul bij 'datum' een geldige datum in ({DATUM_VOORBEELD}; nu "
                      f"'{tekst(datums[0]) if datums else ''}').")
    uit["d"] = d
    # tijden: '10.00 - 10.15' (begin en eind) of alleen een begintijd
    starts, duren = [], []
    for v in waarden.get("tijden", []):
        delen = [naar_minuten(x) for x in re.split(r"\s*[-–]\s*", tekst(v))] if not isinstance(v, (int, float)) else [naar_minuten(v)]
        if not delen or None in delen or len(delen) > 2 or (len(delen) == 2 and delen[1] <= delen[0]):
            fouten.append(f"Blad '{blad}': tijd '{tekst(v)}' is ongeldig (gebruik bijv. 10.00 - 10.15).")
            continue
        if delen[0] in starts:
            meldingen.append(("Waarschuwing", f"Blad '{blad}': tijd '{tekst(v)}' staat dubbel; telt één keer."))
            continue
        starts.append(delen[0])
        if len(delen) == 2:
            duren.append(delen[1] - delen[0])
    if duren:
        if len(set(duren)) > 1:
            meldingen.append(("Waarschuwing", f"Blad '{blad}': niet alle tijden zijn even lang; de toetsduur is {duren[0]} min."))
        s["duur"] = duren[0]
    if waarden.get("pauzes"):
        pz = lees_pauzes(", ".join(tekst(v) for v in waarden["pauzes"]))
        if pz is None:
            fouten.append(f"Blad '{blad}': pauzes moeten zoiets zijn als '11.15-11.40' (van vóór tot, niet overlappend).")
        else:
            s["pauzes"] = pz
    starts = sorted(starts) or tijdgrid(s)
    if starts:
        s["start"], s["eind"] = starts[0], max(starts) + s["duur"]
    kamers = [kenmerken(v) for v in waarden.get("lokalen", [])]
    lokalen = [x[0] for x in kamers]
    lijst = lambda naam: [tekst(v) for v in waarden.get(naam, [])]  # noqa: E731
    casus_k = [kenmerken(v) for v in waarden.get("casus", [])]
    casus_na_k = [kenmerken(v) for v in waarden.get("casus_na", [])]
    casus, casus_na, vaardigheid = [x[0] for x in casus_k], [x[0] for x in casus_na_k], lijst("vaardigheid")
    if len(casus) != len(lokalen):
        meldingen.append(("Waarschuwing", f"Blad '{blad}': {len(lokalen)} lokalen maar {len(casus)} casussen; "
                          "ze worden op volgorde gekoppeld."))
    # beoordelaars: hele dag op volgorde; wie maar een deel van de dag kan, vormt op tijdvolgorde een keten in één
    # lokaal (ochtend -> middag, of 10-12 -> 12-15 -> ...); opmerking tussen haakjes zonder tijd = reserve
    inzet = [lees_inzet(tekst(v), s["middag"]) for v in waarden.get("inzet", [])]
    heel = [p for p in inzet if p["tot"] is None and p["van"] is None and not p["opm"]]
    deels = [p for p in inzet if p["tot"] is not None or p["van"] is not None]
    koppels: list = []
    for p in sorted(deels, key=lambda q: (q["van"] if q["van"] is not None else -1, inzet.index(q))):
        keten = next((k for k in koppels if k[-1]["tot"] is not None and p["van"] is not None and k[-1]["tot"] <= p["van"]), None)
        if keten:
            keten.append(p)
        else:
            koppels.append([p])
    gebruikt = [p for k in koppels for p in k]
    reserve = [p for p in inzet if p not in heel and p not in gebruikt]
    per_lokaal = [[p] for p in heel] + koppels
    while len(per_lokaal) < len(lokalen) and reserve:
        p = reserve.pop(0)
        meldingen.append(("Waarschuwing", f"Reserve {wie_tekst(p)} ingezet omdat er anders te weinig beoordelaars zijn"
                          + (f" (opmerking: {p['opm']})." if p["opm"] else ".")))
        per_lokaal.append([p])
    for p in [q for k in per_lokaal[len(lokalen):] for q in k] + reserve:
        meldingen.append(("Info", f"Niet ingedeeld (reserve): {p['tekst']}."))
    # Casus k (met beoordelaar en vaardigheid) naar een lokaal dat de gevraagde kenmerken heeft, bijv. (keuken).
    eisen = [casus_k[k][1] if k < len(casus_k) else [] for k in range(len(lokalen))]
    plek = koppel_lokalen(kamers, eisen)
    if d is not None:
        meld_koppeling(kamers, [casus[k] if k < len(casus) else "" for k in range(len(lokalen))], eisen, plek, d, meldingen)
    indeling = []
    for k in range(len(lokalen)):
        lokaal = lokalen[plek[k]]
        if k < len(casus_na_k) and casus_na_k[k][1] and not past_bij(casus_na_k[k][1], kamers[plek[k]][1]):
            meldingen.append(("Waarschuwing", f"Casus na pauze {casus_na[k]} vraagt {', '.join(casus_na_k[k][1])}, "
                              f"maar lokaal {lokaal} heeft dat niet."))
        wie = per_lokaal[k] if k < len(per_lokaal) else []
        indeling.append(f"{lokaal} {' → '.join(wie_tekst(p) for p in wie) or '(geen beoordelaar)'}")
        if not wie:
            meldingen.append(("Waarschuwing", f"Lokaal {lokaal} heeft geen beoordelaar; daar wordt niemand ingeroosterd."))
        if d is None:
            continue
        for t in starts:
            uit["slots"].append({"n": len(uit["slots"]), "d": d, "t": t, "lokaal": lokaal,
                                 "casus": casus[k] if k < len(casus) else "",
                                 "casus_na": casus_na[k] if k < len(casus_na) else "",
                                 "vaardigheid": vaardigheid[k] if k < len(vaardigheid) else ""})
            past_ = [p for p in wie if (p["van"] is None or p["van"] <= t) and (p["tot"] is None or t + s["duur"] <= p["tot"])]
            if past_:
                uit["beo"][f"{d}_{t}_{norm(lokaal)}"] = [{"naam": past_[0]["naam"], "casus": "", "lokaal": ""}]
    if not lokalen:
        fouten.append(f"Blad '{blad}': geen lokalen ingevuld.")
    nodig = re.search(r"\d+", " ".join(tekst(v) for v in waarden.get("simus", [])))
    uit["nodig"] = int(nodig[0]) if nodig else 0
    if waarden.get("studenten"):
        rijen = [["Studentnummer", "Voornaam", "Tussenvoegsel", "Achternaam"]]
        for v in waarden["studenten"]:
            w = tekst(v).split()
            j = 1
            while j < len(w) - 1 and w[j].islower():
                j += 1
            rijen.append(["", w[0], " ".join(w[1:j]), " ".join(w[j:])])
        uit["studenten"] = rijen
    if not fouten:
        meldingen.append(("Info", f"Toetsdag uit blad '{blad}' ({datum_tekst(d)}): {len(lokalen)} lokalen, {len(starts)} "
                          f"tijden ({tijd_tekst(s['start'])}-{tijd_tekst(s['eind'])}), toets {s['duur']} min, "
                          f"{len(s['pauzes'])} pauze(s). Tijden en pauzes uit dit blad gaan voor de instellingen."))
        meldingen.append(("Info", "Indeling: " + " · ".join(indeling) + "."))
    return uit


def kandidaten(beschikbaar: dict, slot: dict) -> list:
    """Per lokaal (toetsdag/simu-dagblad) of per tijd (oude bladen)."""
    return beschikbaar.get(f"{slot['d']}_{slot['t']}_{norm(slot['lokaal'])}") or beschikbaar.get(f"{slot['d']}_{slot['t']}", [])


# ---------------------------------------------------------------- plannen

def maak_rooster(bladen: dict, modus: str) -> dict:
    """Berekent het rooster. Geeft {'fouten', 'meldingen', 'bladen', 'namen'} terug.

    'bladen' bevat de nieuwe inhoud van het rooster, de uitdraai, de meldingen en
    (alleen bij eerste kans) het studentenblad; bij fouten alleen de meldingen.
    Ontbreken er rijen in het blad Instellingen, dan wordt het aangevuld.
    Slots komen uit het blad Toetsdag; is dat leeg en bestaat het oude Slots-blad, dan uit dat blad.
    """
    fouten: list[str] = []
    meldingen: list[tuple[str, str]] = []
    s = lees_instellingen(bladen.get(INSTELLINGEN_BLAD), fouten)
    m = MODI[modus]
    namen = {"rooster": s[m["rooster"]], "uitdraai": s[m["uitdraai"]], "meldingen": s["blad_meldingen"]}
    uit = {}
    aangevuld = instellingen_aangevuld(bladen.get(INSTELLINGEN_BLAD))
    if aangevuld:
        uit[INSTELLINGEN_BLAD] = aangevuld
    dag = bladen.get(s[m["dag"]])
    concept = vind_concept(bladen)
    if concept and heeft_data(dag):
        meldingen.append(("Waarschuwing", f"Blad '{s[m['dag']]}' wordt niet gebruikt: het conceptblad '{concept}' gaat voor."))
    c = lees_concept(bladen[concept], concept, s, fouten, meldingen) if concept else {"studenten": None, "nodig": 0}
    studenten_blad = bladen.get(s["blad_studenten"])
    if not heeft_data(studenten_blad) and c["studenten"]:
        studenten_blad = c["studenten"]
        meldingen.append(("Info", f"Studenten uit blad '{concept}' ({len(studenten_blad) - 1}); "
                          f"geen blad '{s['blad_studenten']}' met studenten gevonden."))
    studenten, studenten_kop = lees_studenten(studenten_blad, s["blad_studenten"], fouten, meldingen)
    dagmodus = heeft_data(dag) or bool(concept) or not bladen.get(s[m["slots"]])
    if concept:
        slots, beoordelaars = c["slots"], c["beo"]
    elif dagmodus and not dag:
        fouten.append(f"Blad '{s[m['dag']]}' ontbreekt of is leeg.")
        slots, beoordelaars = [], {}
    elif dagmodus:
        slots, beoordelaars = lees_toetsdag(dag, s[m["dag"]], s, fouten, meldingen)
    else:
        slots = lees_slots(bladen.get(s[m["slots"]]), s[m["slots"]], fouten)
        beoordelaars = lees_beschikbaar(bladen.get(s["blad_beo"]), s["blad_beo"], ("beoordelaar", "naam"), fouten)
    simu_blad = bladen.get(s["blad_simu"])
    simu_dag = None
    if simu_blad and kolom(simu_blad[0], "tijd") < 0:
        simu_dag = lees_simu_dag(simu_blad, s["blad_simu"], fouten, meldingen)
        simus = {}
    else:
        simus = lees_beschikbaar(simu_blad, s["blad_simu"], ("simuclint", "simu", "simulant", "naam"), fouten)
    if fouten:
        rijen = [MELDINGEN_KOP] + [["Fout", f] for f in fouten] + [list(x) for x in meldingen]
        rijen.append(["Info", f"{m['naam']}: rooster NIET aangepast door bovenstaande fouten."])
        uit[namen["meldingen"]] = rijen
        return {"fouten": fouten, "meldingen": meldingen, "bladen": uit, "namen": namen}

    pz = s["pauzes"]
    pauzenr = lambda t: sum(1 for p in pz if p[1] <= t)  # noqa: E731
    regel = s["casus_na"]
    wissel = lambda nr: regel != "Nee" and (nr % 2 == 1 if regel == "Na elke pauze afwisselen" else nr >= int(regel[-1]))  # noqa: E731
    if dagmodus and not concept:
        meldingen.append(("Info", f"Slots gemaakt uit blad '{s[m['dag']]}': {len(slots)} slots "
                          f"({tijd_tekst(s['start'])}-{tijd_tekst(s['eind'])}, elke {s['duur'] + s['wissel']} min)."))
    bruikbaar = []
    in_pauze = 0
    gelijk: list = []
    for slot in slots:
        if any(p[0] <= slot["t"] < p[1] for p in pz):
            in_pauze += 1
            continue
        if slot["casus_na"] and norm(slot["casus_na"]) == norm(slot["casus"]) and f"{slot['d']}_{slot['lokaal']}" not in gelijk:
            gelijk.append(f"{slot['d']}_{slot['lokaal']}")
            meldingen.append(("Waarschuwing", f"Lokaal {slot['lokaal']} op {datum_tekst(slot['d'])}: "
                              f"'Casus na pauze' is gelijk aan 'Casus' ({slot['casus']})."))
        if slot["casus_na"] and wissel(pauzenr(slot["t"])):
            slot["casus"] = slot["casus_na"]
        bruikbaar.append(slot)
    if in_pauze:
        meldingen.append(("Info", f"{in_pauze} slot(s) vallen in een pauze en zijn overgeslagen."))
    if simu_dag is not None and c["nodig"]:
        # 'simu's: 7 x' in het conceptblad: alleen de eerste 7 voor die dag, de rest is reserve.
        die_dag = [x for x in simu_dag if x["d"] is None or x["d"] == c["d"]]
        if len(die_dag) > c["nodig"]:
            meldingen.append(("Info", f"{c['nodig']} simucliënten nodig; reserve: {', '.join(x['naam'] for x in die_dag[c['nodig']:])}."))
            simu_dag = [x for x in simu_dag if x not in die_dag[c["nodig"]:]]
    if simu_dag is not None:
        simus = simu_rotatie(bruikbaar, simu_dag, s, meldingen)

    doel = [st for st in studenten if st["her"] == (modus == HER)]
    volgorde = sorted(doel, key=lambda st: (0 if s["ttv_eerst"] == "Ja" and st["ttv"] else 1 if st["voorkeur"] else 2, st["i"]))
    gewoon = sorted(bruikbaar, key=lambda x: (x["d"], x["t"], x["n"]))
    mid = s["middag"]
    per_voorkeur = {
        "": gewoon, "Vroeg": gewoon,
        "Ochtend": sorted(bruikbaar, key=lambda x: (x["d"], 0 if x["t"] < mid else 1, x["t"], x["n"])),
        "Middag": sorted(bruikbaar, key=lambda x: (x["d"], 0 if x["t"] >= mid else 1, x["t"], x["n"])),
        "Laat": sorted(bruikbaar, key=lambda x: (x["d"], -x["t"], x["n"])),
    }

    def tot_pauze(t):
        # Eind van blok vóór pauze: hoe dichter een slot bij de volgende pauze begint, hoe eerder.
        return min([p[0] - t for p in pz if p[0] > t] or [1000000])
    ttv_volgorde = sorted(bruikbaar, key=lambda x: (x["d"], tot_pauze(x["t"]), -x["t"], x["n"])) \
        if s["ttv_plek"] == "Eind van blok vóór pauze" else gewoon
    deels = s["deelnaam"] == "Ja"
    bezet: dict = {}

    def vrij(sleutel, a, b):
        return all(b <= x or a >= y for x, y in bezet.get(sleutel, []))

    def anders(regel_her, strikt):
        # Moet het anders dan bij de 1e kans? Verplicht altijd, Voorkeur alleen in de eerste (strikte) ronde.
        return modus == HER and (s[regel_her] == "Verplicht" or (strikt and s[regel_her] == "Voorkeur"))

    def kies(lijst, slot, a, b, soort, st, strikt):
        for kandidaat in lijst:
            if kandidaat["casus"] and norm(kandidaat["casus"]) != norm(slot["casus"]):
                continue
            if kandidaat["lokaal"] and norm(kandidaat["lokaal"]) != norm(slot["lokaal"]):
                continue
            if soort == "b" and beperkt(kandidaat["naam"], st["verboden"], deels):
                continue
            if soort == "b" and anders("her_beo", strikt) and st["eb"] and norm(kandidaat["naam"]) == norm(st["eb"]):
                continue
            if vrij((soort, slot["d"], kandidaat["naam"]), a, b):
                return kandidaat["naam"]
        return None

    def past(slot, a, b, ttv):
        # Gewone toets niet in een pauze; TTV mag doorlopen in de pauze direct erna. Toetsdag: niet na de eindtijd.
        if any(a < p[0] and b > (p[1] if ttv else p[0]) for p in pz):
            return False
        return not (dagmodus and b > s["eind"]) and vrij(("l", slot["d"], slot["lokaal"]), a, b)

    def boek(slot, a, b, st, strikt):
        beo = kies(kandidaten(beoordelaars, slot), slot, a, b, "b", st, strikt)
        simu = beo and kies(kandidaten(simus, slot), slot, a, b, "s", st, strikt)
        if not simu:
            return None
        for k in (("l", slot["d"], slot["lokaal"]), ("b", slot["d"], beo), ("s", slot["d"], simu)):
            bezet.setdefault(k, []).append((a, b))
        return {"st": st, "slot": slot, "eind": b, "beo": beo, "simu": simu}

    aantal: dict = {}  # studenten per lokaal (per dag) tot nu toe

    def geordend(lijst):
        # 'Evenredig': binnen hetzelfde tijdstip eerst het lokaal met de minste studenten; tijdvolgorde blijft gelijk.
        if s["verdelen"] != "Evenredig":
            return lijst
        groep: dict = {}
        for x in lijst:
            groep.setdefault((x["d"], x["t"]), len(groep))
        return sorted(lijst, key=lambda x: (groep[(x["d"], x["t"])], aantal.get((x["d"], norm(x["lokaal"])), 0), x["n"]))

    def probeer(st, strikt):
        duur = s["ttv"] if st["ttv"] else s["duur"]
        for slot in geordend(ttv_volgorde if st["ttv"] else per_voorkeur[st["voorkeur"]]):
            a, b = slot["t"], slot["t"] + duur
            if not past(slot, a, b, st["ttv"]):
                continue
            if anders("her_casus", strikt) and st["ec"] and norm(slot["casus"]) == norm(st["ec"]):
                continue
            r = boek(slot, a, b, st, strikt)
            if r:
                aantal[(slot["d"], norm(slot["lokaal"]))] = aantal.get((slot["d"], norm(slot["lokaal"])), 0) + 1
                return r
        return None

    toegewezen: dict = {}
    for strikt in ([True, False] if modus == HER else [False]):
        for st in volgorde:
            if st["i"] not in toegewezen:
                r = probeer(st, strikt)
                if r:
                    toegewezen[st["i"]] = r

    vrije = []
    if s["vrije_plekken"] == "Ja":
        niemand = {"verboden": [], "eb": "", "ec": "", "ttv": False}
        for slot in gewoon:
            a, b = slot["t"], slot["t"] + s["duur"]
            r = boek(slot, a, b, niemand, False) if past(slot, a, b, False) else None
            if r:
                r["st"] = None
                vrije.append(r)

    # Lokaal waar niemand terechtkan: zeg welk lokaal en waarom (anders 'verdwijnt' bijv. een beoordelaar zonder uitleg).
    gebruikt = {(r["slot"]["d"], norm(r["slot"]["lokaal"])) for r in list(toegewezen.values()) + vrije}
    niemand = {"verboden": [], "eb": "", "ec": "", "ttv": False}
    for d, lok in dict.fromkeys((x["d"], x["lokaal"]) for x in gewoon):
        if (d, norm(lok)) in gebruikt:
            continue
        rs = [x for x in gewoon if x["d"] == d and x["lokaal"] == lok]
        beo_namen = list(dict.fromkeys(p_["naam"] for x in rs for p_ in kandidaten(beoordelaars, x)))
        bruik = [x for x in rs if past(x, x["t"], x["t"] + s["duur"], False)
                 and kies(kandidaten(beoordelaars, x), x, x["t"], x["t"] + s["duur"], "b", niemand, False)
                 and kies(kandidaten(simus, x), x, x["t"], x["t"] + s["duur"], "s", niemand, False)]
        if bruik:
            continue  # lokaal was bruikbaar maar niet nodig
        reden = ("er is op geen enkel moment een beoordelaar" if not beo_namen else
                 "er is op geen enkel moment een simucliënt die daar kan" if not any(kandidaten(simus, x) for x in rs) else
                 "de beoordelaar of de simucliënten zijn steeds al in een ander lokaal bezig (bijv. minder simu's dan lokalen; "
                 "zet meer simu's in of haal een lokaal weg)")
        meldingen.append(("Waarschuwing", f"Lokaal {lok} op {datum_tekst(d)} ({' → '.join(beo_namen) or 'geen beoordelaar'}) "
                          f"krijgt niemand: {reden}."))

    if s["rooster_sort"] == "Lokaal, dan tijd":
        rooster_sleutel = lambda r: (r["slot"]["d"], r["slot"]["lokaal"].lower(), r["slot"]["t"], r["slot"]["n"])  # noqa: E731
    else:
        rooster_sleutel = lambda r: (r["slot"]["d"], r["slot"]["t"], r["slot"]["lokaal"].lower(), r["slot"]["n"])  # noqa: E731
    if s["uitdraai_sort"] == "Achternaam":
        uitdraai_sleutel = lambda r: (r["st"]["achternaam"].lower(), r["st"]["voornaam"].lower(), r["st"]["i"])  # noqa: E731
    elif s["uitdraai_sort"] == "Tijd":
        uitdraai_sleutel = lambda r: (r["slot"]["d"], r["slot"]["t"], r["st"]["achternaam"].lower(), r["st"]["i"])  # noqa: E731
    else:
        uitdraai_sleutel = lambda r: (r["slot"]["d"], r["slot"]["lokaal"].lower(), r["slot"]["t"], r["st"]["i"])  # noqa: E731

    rooster = [ROOSTER_KOP]
    uitdraai = [UITDRAAI_KOP]
    for r in sorted(list(toegewezen.values()) + vrije, key=rooster_sleutel):
        st, slot = r["st"], r["slot"]
        if st is None:
            rooster.append([slot["d"], slot["t"] / 1440, slot["lokaal"], slot["casus"], slot["vaardigheid"], r["beo"],
                            r["simu"], "", "", "", "", "", VRIJ, r["eind"] / 1440])
            continue
        naam = f"{st['voornaam']} {st['achternaam']}"
        bijz = [f"TTV ({s['ttv']} min)"] if st["ttv"] else []
        if modus == HER and st["eb"] and norm(st["eb"]) == norm(r["beo"]):
            bijz.append("Zelfde beoordelaar als 1e kans")
            meldingen.append(("Waarschuwing", f"{naam}: zelfde beoordelaar als 1e kans."))
        if modus == HER and st["ec"] and norm(st["ec"]) == norm(slot["casus"]):
            bijz.append("Zelfde casus als 1e kans")
            meldingen.append(("Waarschuwing", f"{naam}: zelfde casus als 1e kans."))
        if (st["voorkeur"] == "Ochtend" and slot["t"] >= mid) or (st["voorkeur"] == "Middag" and slot["t"] < mid):
            meldingen.append(("Waarschuwing", f"{naam}: voorkeur {st['voorkeur'].lower()} niet gehaald "
                              f"({tijd_tekst(slot['t'])})."))
        if st["opm"]:
            bijz.append(st["opm"])
        rooster.append([slot["d"], slot["t"] / 1440, slot["lokaal"], slot["casus"], slot["vaardigheid"], r["beo"], r["simu"],
                        st["nr"], st["voornaam"], st["tv"], st["achternaam"], st["klas"], "; ".join(bijz), r["eind"] / 1440])
    for r in sorted(toegewezen.values(), key=uitdraai_sleutel):
        st = r["st"]
        naam = " ".join(x for x in (st["voornaam"], st["tv"], st["achternaam"]) if x)
        tijd = tijd_tekst(r["slot"]["t"])
        uitdraai.append([naam, r["slot"]["lokaal"], f"{datum_tekst(r['slot']['d'])} {tijd}" if s["uitdraai_datum"] == "Ja" else tijd])
    niet = [st for st in volgorde if st["i"] not in toegewezen]
    for st in sorted(niet, key=lambda x: x["i"]):
        rooster.append(["", "", "", "", "", "", "", st["nr"], st["voornaam"], st["tv"], st["achternaam"], st["klas"], NIET, ""])
        meldingen.append(("Waarschuwing", f"{st['voornaam']} {st['achternaam']} is NIET ingeroosterd: geen vrij lokaal, "
                          "beoordelaar of simucliënt gevonden (let ook op uitsluitingen en verplichte herkansingsregels)."))
    if vrije:
        meldingen.append(("Info", f"{len(vrije)} vrije plek(ken) over. Ze staan als '{VRIJ}' in het rooster, "
                          "niet op de uitdraai of in de PDF's."))

    uit[namen["rooster"]] = rooster
    uit[namen["uitdraai"]] = uitdraai
    if modus == KANS1:
        kop = list(studenten_kop)
        cols = []
        for naam in EERSTE_KOP:
            c = kolom(kop, norm(naam))
            if c < 0:
                c = len(kop)
                kop.append(naam)
            cols.append(c)
        nieuw = [list(r) + [""] * (len(kop) - len(r)) for r in studenten_blad]
        nieuw[0] = kop
        for st in doel:
            r = toegewezen.get(st["i"])
            waarden = [r["slot"]["d"], r["slot"]["t"] / 1440, r["beo"], r["slot"]["casus"]] if r else ["", "", "", ""]
            for c, w in zip(cols, waarden):
                nieuw[st["i"]][c] = w
        uit[s["blad_studenten"]] = nieuw
    info = f"{m['naam']}: {len(toegewezen)} van {len(doel)} studenten ingeroosterd (jaar {s['jaar']}, toets {s['duur']} min, TTV {s['ttv']} min"
    info += f", pauzes {pauzes_tekst(pz)})." if pz else ", geen pauze)."
    meldingen.insert(0, ("Info", info))
    uit[namen["meldingen"]] = [MELDINGEN_KOP] + [list(x) for x in meldingen]
    return {"fouten": [], "meldingen": meldingen, "bladen": uit, "namen": namen}


# ---------------------------------------------------------------- Excel-bestanden (desktop-app)

OPMAAK = {"datum": "dd-mm-yyyy", "eerstedatum": "dd-mm-yyyy", "tijd": "hh:mm", "eindtijd": "hh:mm", "eerstetijd": "hh:mm"}


def excel_waarde(v):
    """openpyxl-waarde -> waarde zoals Excel Online getValues() hem geeft."""
    if isinstance(v, dt.datetime):
        x = v - dt.datetime(1899, 12, 30)
        return x.days + x.seconds / 86400 if x.seconds else x.days
    if isinstance(v, dt.date):
        return (v - dt.date(1899, 12, 30)).days
    if isinstance(v, dt.time):
        return (v.hour * 3600 + v.minute * 60 + v.second) / 86400
    if isinstance(v, dt.timedelta):
        return v.total_seconds() / 86400
    return "" if v is None else v


def lees_werkmap(pad) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(pad, data_only=True)
    bladen = {}
    for ws in wb.worksheets:
        rijen = [[excel_waarde(c) for c in rij] for rij in ws.iter_rows(values_only=True)]
        while rijen and leeg(rijen[-1]):
            rijen.pop()
        bladen[ws.title] = rijen
    return bladen


def is_studentenblad(rijen: list) -> bool:
    """Het studentenblad herken je aan de kolom EersteDatum; daarvan worden alleen de Eerste*-kolommen geschreven."""
    return bool(rijen) and any(norm(k) == "eerstedatum" for k in rijen[0])


def schrijf_bladen(wb, bladen: dict) -> None:
    """Schrijft matrices in een openpyxl-werkmap (bestaande bladen worden overschreven)."""
    from openpyxl.styles import Font
    tijden = {norm(i["label"]) for i in INSTELLINGEN if i["soort"] == "tijd"}
    for naam, rijen in bladen.items():
        if is_studentenblad(rijen) and naam in wb.sheetnames:
            # Andere cellen (en formules) blijven ongemoeid.
            ws = wb[naam]
            eerste = {norm(k) for k in EERSTE_KOP}
            for c, titel in enumerate(rijen[0], start=1):
                if norm(titel) in eerste:
                    for r, rij in enumerate(rijen, start=1):
                        ws.cell(r, c).value = None if rij[c - 1] == "" else rij[c - 1]
                        if r > 1 and norm(titel) in OPMAAK:
                            ws.cell(r, c).number_format = OPMAAK[norm(titel)]
            continue
        if naam in wb.sheetnames:
            ws = wb[naam]
            ws.delete_rows(1, ws.max_row)
        else:
            ws = wb.create_sheet(naam)
        for rij in rijen:
            ws.append([None if w == "" else w for w in rij])
        kop = rijen[0] if rijen else []
        for c, titel in enumerate(kop, start=1):
            ws.cell(1, c).font = Font(bold=True)
            fmt = OPMAAK.get(norm(titel))
            if fmt and kop != MELDINGEN_KOP:
                for r in range(2, len(rijen) + 1):
                    ws.cell(r, c).number_format = fmt
            breedte = max((len(tekst(rij[c - 1])) for rij in rijen if c - 1 < len(rij)), default=8)
            ws.column_dimensions[ws.cell(1, c).column_letter].width = min(max(breedte, 8) + 2, 90)
        if list(kop) == INSTELLINGEN_KOP:
            for r in range(2, len(rijen) + 1):
                if norm(cel(rijen[r - 1], 0)) in tijden:
                    ws.cell(r, 2).number_format = "hh:mm"
