"""Uitvoer van de app: Excel-werkmap, PDF's om te printen/publiceren en een leeg sjabloon."""
from __future__ import annotations

import datetime as dt
from html import escape

from PySide6.QtCore import QMarginsF
from PySide6.QtGui import QFont, QPageLayout, QPageSize, QPdfWriter, QTextDocument

import rooster_core as rc

DAGEN = ["maandag", "dinsdag", "woensdag", "donderdag", "vrijdag", "zaterdag", "zondag"]
STUDENTEN_KOP = ["Studentnummer", "Voornaam", "Tussenvoegsel", "Achternaam", "klas", "Niet inroosteren bij SB-docent",
                 "Niet inroosteren bij de vaardighedendocent", "TTV", "Voorkeur", "Opmerkingen", "IsHerkanser"] + rc.EERSTE_KOP
CSS = """
body { font-family: Arial, Helvetica, sans-serif; color: #0a0a0a; }
h1 { font-size: 18pt; font-weight: 600; margin: 0 0 2pt 0; }
p.sub { color: #737373; font-size: 10pt; margin: 0 0 12pt 0; }
table { border-collapse: collapse; }
th { background: #f5f5f5; color: #404040; font-weight: 600; text-align: left; border-bottom: 1px solid #d4d4d4; }
td { border-bottom: 1px solid #e5e5e5; }
"""


def dag_tekst(serie: int) -> str:
    d = dt.date(1899, 12, 30) + dt.timedelta(days=serie)
    return f"{DAGEN[d.weekday()]} {rc.datum_tekst(serie)}"


def _tabel(kop: list, rijen: list) -> str:
    th = "".join(f"<th>{escape(str(k))}</th>" for k in kop)
    tr = "".join("<tr>" + "".join(f"<td>{escape(rc.tekst(c))}</td>" for c in rij) + "</tr>" for rij in rijen)
    return f"<table width='100%' cellpadding='6' cellspacing='0'><thead><tr>{th}</tr></thead>{tr}</table>"


def html_uitdraai(titel: str, matrix: list) -> str:
    """Voor studenten: alleen naam, lokaal en tijd."""
    return (f"<style>{CSS}</style><h1>{escape(titel)}</h1>"
            f"<p class='sub'>{len(matrix) - 1} studenten</p>{_tabel(matrix[0], matrix[1:])}")


def html_rooster(titel: str, matrix: list) -> str:
    """Voor docenten: per dag en lokaal een eigen pagina (om bij het lokaal op te hangen)."""
    kop = matrix[0]
    i = {k: kop.index(k) for k in kop}
    groepen: dict = {}
    niet = []
    for rij in matrix[1:]:
        if rij[i["bijzheden"]] == rc.VRIJ:
            continue  # vrije plekken alleen in het Excel-rooster
        if rij[i["bijzheden"]] == rc.NIET:
            niet.append(rij)
        else:
            groepen.setdefault((rij[i["Datum"]], rij[i["Lokaal"]]), []).append(rij)
    delen = []
    for n, ((datum, lokaal), rijen) in enumerate(sorted(groepen.items(), key=lambda g: (g[0][0], rc.tekst(g[0][1]).lower()))):
        rijen.sort(key=lambda r: r[i["Tijd"]])
        casussen = sorted({rc.tekst(r[i["Casus"]]) for r in rijen if rc.tekst(r[i["Casus"]])})
        regels = [[f"{rc.tijd_tekst(rc.naar_minuten(r[i['Tijd']]))}–{rc.tijd_tekst(rc.naar_minuten(r[i['Eindtijd']]))}",
                   " ".join(rc.tekst(r[i[k]]) for k in ("Voornaam", "vv", "Achternaam") if rc.tekst(r[i[k]])),
                   r[i["stnr"]], r[i["klas"]], r[i["Casus"]], r[i["Beoordelaar"]], r[i["simucliënt"]], r[i["bijzheden"]]]
                  for r in rijen]
        breuk = " style='page-break-before: always'" if n else ""
        delen.append(f"<h1{breuk}>Lokaal {escape(rc.tekst(lokaal))} · {escape(dag_tekst(datum))}</h1>"
                     f"<p class='sub'>{escape(titel)} · {len(rijen)} studenten · casus {escape(', '.join(casussen) or '–')}</p>"
                     + _tabel(["Tijd", "Student", "Nr.", "Klas", "Casus", "Beoordelaar", "Simucliënt", "Bijzonderheden"], regels))
    if niet:
        delen.append("<h1 style='page-break-before: always'>Niet ingeroosterd</h1><p class='sub'>Zie het blad Meldingen.</p>"
                     + _tabel(["Student", "Nr.", "Klas"], [[" ".join(rc.tekst(r[i[k]]) for k in ("Voornaam", "vv", "Achternaam")
                                                                   if rc.tekst(r[i[k]])), r[i["stnr"]], r[i["klas"]]] for r in niet]))
    return f"<style>{CSS}</style>" + "".join(delen)


def bewaar_pdf(html: str, pad: str, liggend: bool = False) -> None:
    pdf = QPdfWriter(pad)
    pdf.setPageSize(QPageSize(QPageSize.A4))
    pdf.setPageOrientation(QPageLayout.Landscape if liggend else QPageLayout.Portrait)
    pdf.setResolution(300)
    pdf.setPageMargins(QMarginsF(15, 15, 15, 15), QPageLayout.Millimeter)
    pdf.setTitle("Toetsrooster")
    doc = QTextDocument()
    doc.documentLayout().setPaintDevice(pdf)  # punten op PDF-resolutie, niet schermresolutie
    doc.setDefaultFont(QFont("Arial", 10))
    doc.setHtml(html)
    doc.setPageSize(pdf.pageLayout().paintRectPixels(pdf.resolution()).size().toSizeF())
    doc.print_(pdf)


def bewaar_werkmap(bron: str, doel: str, bladen: dict) -> None:
    """Kopie van de bronwerkmap met de nieuwe bladen erin (het origineel blijft onaangeroerd)."""
    from openpyxl import load_workbook
    wb = load_workbook(bron)
    rc.schrijf_bladen(wb, bladen)
    wb.save(doel)


def maak_sjabloon(pad: str) -> None:
    """Lege werkmap met de juiste bladen en kolommen, om zelf in te vullen."""
    from openpyxl import Workbook
    s = rc.STANDAARD
    dag = ["Datum", "Lokaal", "Casus", "Casus na pauze", "Vaardigheid", "Beoordelaar", "Van", "Tot",
           "Beoordelaar 2", "Van 2", "Tot 2", "Beoordelaar 3", "Van 3", "Tot 3"]
    wb = Workbook()
    wb.remove(wb.active)
    rc.schrijf_bladen(wb, {
        s["blad_studenten"]: [STUDENTEN_KOP],
        s["blad_dag1"]: [dag],
        s["blad_dag_her"]: [dag],
        s["blad_simu"]: [["Datum", "simucliënt", "Vanaf", "Tot", "Vast lokaal", "Kan spelen"]],
        rc.INSTELLINGEN_BLAD: rc.instellingen_matrix(),
    })
    wb.save(pad)


# ---------------------------------------------------------------- beschikbaarheid simu's

MAANDEN = {"jan": 1, "feb": 2, "mrt": 3, "maa": 3, "apr": 4, "mei": 5, "jun": 6, "jul": 7, "aug": 8, "sep": 9,
           "okt": 10, "nov": 11, "dec": 12}


def _dag_maand(titel: str):
    """'Fysio jr 2\\n in les\\n22 sept' -> (22, 9); None als er geen datum in staat."""
    import re
    m = re.findall(r"(\d{1,2})\s*([a-z]{3})[a-z]*", titel.lower())
    for dag, maand in reversed(m):
        if maand in MAANDEN and 1 <= int(dag) <= 31:
            return int(dag), MAANDEN[maand]
    return None


def datum_bij(dag_maand, rond: int) -> int | None:
    """Kolomdatum zonder jaartal -> Excel-serienummer, met het jaar het dichtst bij datum `rond`."""
    if not dag_maand:
        return None
    jaar = (dt.date(1899, 12, 30) + dt.timedelta(days=rond)).year
    opties = [rc.serie(j, dag_maand[1], dag_maand[0]) for j in (jaar - 1, jaar, jaar + 1)]
    opties = [o for o in opties if o is not None]
    return min(opties, key=lambda o: abs(o - rond)) if opties else None


def lees_cel_beschikbaar(v, middag: int) -> dict | None:
    """Cel uit het beschikbaarheidsoverzicht -> {'van', 'tot', 'zeker', 'tekst'}; None = niet beschikbaar."""
    import re
    t = rc.tekst(v)
    if not t:
        return None
    if t == "?":
        return {"van": None, "tot": None, "zeker": False, "tekst": t}
    tijd = r"(\d{1,2}(?:[:.]\d{2})?)"
    los = lambda x: rc.naar_minuten(x if re.search(r"[:.]", x) else f"{x}:00")  # noqa: E731
    m = re.search(tijd + r"\s*[-–]\s*" + tijd, t)
    if m:
        return {"van": los(m[1]), "tot": los(m[2]), "zeker": True, "tekst": t}
    m = re.search(r"\btot\s+" + tijd, t.lower())
    if m:
        return {"van": None, "tot": los(m[1]), "zeker": True, "tekst": t}
    m = re.search(r"\b(?:va|vanaf|van)\s+" + tijd, t.lower())
    if m:
        return {"van": los(m[1]), "tot": None, "zeker": True, "tekst": t}
    if "middag" in t.lower():
        return {"van": middag, "tot": None, "zeker": True, "tekst": t}
    return {"van": None, "tot": None, "zeker": True, "tekst": t}  # X of andere aantekening


def lees_beschikbaarheid(pad: str, middag: int = 12 * 60) -> list:
    """Overzicht 'Beschikbaarheid simu's': rij 1 = namen + kolomtitels met datum, rij 2 = aantal nodig,
    rij 3 = tijden, daarna één rij per simu. Geeft per datumkolom {titel, dag_maand, nodig, simus}."""
    import re
    from openpyxl import load_workbook
    ws = load_workbook(pad, data_only=True).worksheets[0]
    rijen = [list(r) for r in ws.iter_rows(values_only=True)]
    if len(rijen) < 4:
        raise ValueError("Te weinig rijen: verwacht kolomtitels, 'nodig', tijden en daarna één rij per simucliënt.")
    kolommen = []
    for c in range(1, len(rijen[0])):
        titel = rc.tekst(rijen[0][c])
        if not titel:
            continue
        nodig = re.search(r"\d+", rc.tekst(rc.cel(rijen[1], c)))
        simus = []
        for rij in rijen[3:]:
            naam = rc.tekst(rc.cel(rij, 0))
            b = lees_cel_beschikbaar(rc.cel(rij, c), middag) if naam else None
            if b:
                simus.append({"naam": naam, **b})
        kolommen.append({"titel": " ".join(titel.split()), "dag_maand": _dag_maand(titel),
                         "nodig": int(nodig[0]) if nodig else None, "tijd": rc.tekst(rc.cel(rijen[2], c)), "simus": simus})
    return kolommen


def simu_blad_met(oud: list | None, datum: int, gekozen: list) -> list:
    """Blad Simu (dagvorm) met de rijen voor `datum` vervangen door `gekozen` ({naam, van, tot}).
    Al ingevulde 'Vast lokaal' en 'Kan spelen' blijven per simu bewaard. Een oud blad met kolom Tijd wordt vervangen."""
    kop = ["Datum", "simucliënt", "Vanaf", "Tot", "Vast lokaal", "Kan spelen"]
    rijen, vast = [], {}
    if oud and rc.kolom(oud[0], "tijd") < 0:
        k = [rc.kolom(oud[0], "datum"), rc.kolom(oud[0], "simuclint", "simu", "simulant", "naam"),
             rc.kolom(oud[0], "vanaf", "van"), rc.kolom(oud[0], "tot"), rc.kolom(oud[0], "vastlokaal"),
             rc.kolom(oud[0], "kanspelen")]
        for r in oud[1:]:
            if rc.leeg(r):
                continue
            if rc.naar_datum(rc.cel(r, k[0])) == datum:
                vast[rc.norm(rc.cel(r, k[1]))] = [rc.cel(r, k[4]), rc.cel(r, k[5])]
            else:
                rijen.append([rc.cel(r, i) for i in k])
    tijd = lambda m: "" if m is None else m / 1440  # noqa: E731
    return [kop] + rijen + [[datum, s["naam"], tijd(s["van"]), tijd(s["tot"])] + vast.get(rc.norm(s["naam"]), ["", ""])
                            for s in gekozen]


def simu_vast_bijwerken(oud: list, waarden: dict) -> list:
    """Blad Simu (dagvorm) met per rijnummer {rij: (vast lokaal, kan spelen)} ingevuld; kolommen worden zo nodig toegevoegd."""
    kop = list(oud[0])
    cols = []
    for naam, titel in (("vastlokaal", "Vast lokaal"), ("kanspelen", "Kan spelen")):
        c = rc.kolom(kop, naam)
        if c < 0:
            c = len(kop)
            kop.append(titel)
        cols.append(c)
    nieuw = [kop] + [list(r) + [""] * (len(kop) - len(r)) for r in oud[1:]]
    for i, w in waarden.items():
        nieuw[i][cols[0]], nieuw[i][cols[1]] = w
    return nieuw
