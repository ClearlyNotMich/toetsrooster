"""Rooktest van de Qt-app zonder scherm: openen, rooster maken, opslaan als Excel en PDF."""
import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import load_workbook  # noqa: E402

import maak_voorbeeld as mv  # noqa: E402
import rooster_core as rc  # noqa: E402
import roosterapp  # noqa: E402
import uitvoer  # noqa: E402


def test_app_maakt_en_bewaart(tmp_path):
    bron = tmp_path / "werkmap.xlsx"
    mv.schrijf(bron)
    assert roosterapp.controle(str(bron), str(tmp_path)) == 0
    uit = rc.lees_werkmap(tmp_path / "controle-Rooster_Kans1.xlsx")
    verwacht = rc.maak_rooster(rc.lees_werkmap(bron), "kans1")["bladen"]
    for blad in ("Rooster_Kans1", "Uitdraai_Kans1"):
        assert [[round(c, 6) if isinstance(c, float) else c for c in r] for r in uit[blad]] == \
               [[round(c, 6) if isinstance(c, float) else ("" if c is None else c) for c in r] for r in verwacht[blad]]
    # Studenten: alleen Eerste*-kolommen veranderd, rest identiek
    oud, nieuw = rc.lees_werkmap(bron)["Studenten"], uit["Studenten"]
    assert [r[:9] for r in oud] == [r[:9] for r in nieuw]
    assert load_workbook(tmp_path / "controle-Rooster_Kans1.xlsx")["Rooster_Kans1"]["B2"].number_format == "hh:mm"
    for naam in ("controle-Rooster_Kans1.pdf", "controle-Rooster_Her.pdf", "controle-Rooster_Kans1-docenten.pdf"):
        assert (tmp_path / naam).read_bytes()[:4] == b"%PDF"
    sjabloon = rc.lees_werkmap(tmp_path / "controle-sjabloon.xlsx")
    assert set(sjabloon) == {"Studenten", "Toetsdag_Kans1", "Toetsdag_Her", "Simu", "Instellingen"}
    assert rc.maak_rooster(sjabloon, "kans1")["fouten"] == []  # leeg maar geldig
    her = rc.lees_werkmap(tmp_path / "controle-Rooster_Her.xlsx")["Rooster_Her"]
    assert sum(1 for rij in her[1:] if rc.VRIJ not in rij) == 2


def test_origineel_wordt_niet_overschreven(tmp_path, monkeypatch):
    bron = tmp_path / "werkmap.xlsx"
    mv.schrijf(bron)
    app = roosterapp.QApplication.instance() or roosterapp.QApplication([])  # noqa: F841
    v = roosterapp.Venster()
    v.open(str(bron))
    v.maak()
    voor = bron.read_bytes()
    monkeypatch.setattr(roosterapp.QMessageBox, "warning", lambda *a: None)
    assert v.opslaan(str(bron)) is None and bron.read_bytes() == voor


def test_fouten_blokkeren_opslaan(tmp_path):
    b = mv.voorbeeld_bladen()
    b["Slots_Kans1"][2][1] = "??"
    from openpyxl import Workbook
    wb = Workbook(); wb.remove(wb.active); rc.schrijf_bladen(wb, b); wb.save(tmp_path / "k.xlsx")
    app = roosterapp.QApplication.instance() or roosterapp.QApplication([])  # noqa: F841
    v = roosterapp.Venster()
    v.open(str(tmp_path / "k.xlsx"))
    r = v.maak()
    assert r["fouten"] and not v.knop_opslaan.isEnabled() and v.tabs.currentWidget() is v.tabel_meldingen


def test_wijziging_na_maken_vraagt_opnieuw_maken(tmp_path):
    bron = tmp_path / "werkmap.xlsx"
    mv.schrijf(bron)
    app = roosterapp.QApplication.instance() or roosterapp.QApplication([])  # noqa: F841
    v = roosterapp.Venster()
    v.open(str(bron))
    v.maak()
    assert v.knop_pdf.isEnabled()
    v.instellingen.velden["duur1"].setValue(20)  # een instelling wijzigen: eerst opnieuw maken
    assert not v.knop_pdf.isEnabled() and not v.knop_opslaan.isEnabled()
    v.modus_groep.buttons()[1].setChecked(True)  # andere toets kiezen: rooster wordt meteen gemaakt
    assert v.resultaat["modus"] == "her" and v.knop_pdf.isEnabled()


def test_instellingen_formulier_heen_en_terug():
    app = roosterapp.QApplication.instance() or roosterapp.QApplication([])  # noqa: F841
    v = roosterapp.Venster()
    assert v.instellingen.waarden() == rc.STANDAARD
    anders = {**rc.STANDAARD, "jaar": 2, "duur2": 20, "pauzes": "12:00-12:30", "start": "08:30", "simu_rotatie": "Nee",
              "ttv_eerst": "Nee", "her_beo": "Verplicht", "blad_meldingen": "Log"}
    v.instellingen.zet(anders)
    assert v.instellingen.waarden() == anders
    v.instellingen.herstel()
    assert v.instellingen.waarden() == rc.STANDAARD


def test_instellingen_komen_in_rooster_en_werkmap(tmp_path):
    bron = tmp_path / "werkmap.xlsx"
    mv.schrijf(bron)
    app = roosterapp.QApplication.instance() or roosterapp.QApplication([])  # noqa: F841
    v = roosterapp.Venster()
    v.open(str(bron))
    v.instellingen.velden["duur1"].setValue(20)
    v.instellingen.velden["blad_rooster1"].setText("Rooster docenten")
    r = v.maak()
    assert "toets 20 min" in r["meldingen"][0][1] and r["namen"]["rooster"] == "Rooster docenten"
    doel = v.opslaan(str(tmp_path / "uit.xlsx"))
    uit = rc.lees_werkmap(doel)
    assert "Rooster docenten" in uit and rc.instellingen_waarden(uit["Instellingen"])["duur1"] == 20


def beschikbaarheid(pad):
    """Verzonnen overzicht in dezelfde vorm als 'Beschikbaarheid simu's'."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["naam", "Fysio jr 2\n in les\n22 sept", "M1.6 jr 1\n14 april", "M1.6 jr 1 HER\n27 mei"])
    ws.append([None, 2, 3, "zelfde 2"])
    ws.append([None, None, "08.30-17.00", None])
    for rij in (["Siem", "X", "X", None], ["Tom", None, "?", "X"], ["Uma", "X", "8.30- 12.15", "X"],
                ["Vera", None, "va 13.45", None], ["Wout", None, "X", None], ["Xan", None, "middag", "tot 12 uur"]):
        ws.append(rij)
    wb.save(pad)


def test_beschikbaarheid_lezen(tmp_path):
    beschikbaarheid(tmp_path / "b.xlsx")
    k = uitvoer.lees_beschikbaarheid(str(tmp_path / "b.xlsx"))
    assert [x["dag_maand"] for x in k] == [(22, 9), (14, 4), (27, 5)] and [x["nodig"] for x in k] == [2, 3, 2]
    april = {s["naam"]: (s["van"], s["tot"], s["zeker"]) for s in k[1]["simus"]}
    assert april == {"Siem": (None, None, True), "Tom": (None, None, False), "Uma": (510, 735, True),
                     "Vera": (825, None, True), "Wout": (None, None, True), "Xan": (720, None, True)}
    assert rc.datum_tekst(uitvoer.datum_bij((14, 4), mv.TOETS)) == "14-04-2026"


def test_simus_importeren_en_opslaan(tmp_path):
    bron = tmp_path / "werkmap.xlsx"
    mv.schrijf(bron)
    beschikbaarheid(tmp_path / "b.xlsx")
    app = roosterapp.QApplication.instance() or roosterapp.QApplication([])  # noqa: F841
    v = roosterapp.Venster()
    v.open(str(bron))
    gezien = {}

    def kies(dlg):
        gezien["kolom"] = dlg.kolom.currentText()
        gezien["vink"] = [s["naam"] for s in dlg.gekozen()]
        return True
    assert v.importeer_simus(str(tmp_path / "b.xlsx"), kies)
    assert gezien == {"kolom": "M1.6 jr 1 14 april", "vink": ["Siem", "Uma", "Vera"]}  # '?' (Tom) niet voorgevinkt
    simu = v.bladen["Simu"]
    assert [r[1] for r in simu[1:] if r[0] == mv.TOETS] == ["Siem", "Uma", "Vera"]
    assert [r[1] for r in simu[1:] if r[0] == mv.HER] == ["Siem", "Tom", "Uma"]  # andere dag blijft staan
    r = v.maak()
    assert r["fouten"] == [] and any("11:45-13:00: 2 simucliënt(en) beschikbaar voor 5 lokalen" in m[1] for m in r["meldingen"])  # Siem + Uma (tot 12:15)
    uit = rc.lees_werkmap(v.opslaan(str(tmp_path / "uit.xlsx")))
    assert [r[1] for r in uit["Simu"][1:] if r[0] == mv.TOETS] == ["Siem", "Uma", "Vera"]
    assert rc.naar_minuten(next(r for r in uit["Simu"][1:] if r[1] == "Uma" and r[0] == mv.TOETS)[3]) == 735


def test_vrije_plekken_niet_in_pdf(tmp_path):
    bron = tmp_path / "werkmap.xlsx"
    mv.schrijf(bron)
    app = roosterapp.QApplication.instance() or roosterapp.QApplication([])  # noqa: F841
    v = roosterapp.Venster()
    v.open(str(bron))
    r = v.maak()
    rooster = r["bladen"]["Rooster_Kans1"]
    assert sum(rc.VRIJ in rij for rij in rooster) == 24 and v.stat_in.waarde.text() == "61 / 61"
    html = uitvoer.html_rooster("t", rooster) + uitvoer.html_uitdraai("t", r["bladen"]["Uitdraai_Kans1"])
    assert rc.VRIJ not in html and html.count("<tr>") == 61 + 61 + 5 + 1  # 5 lokaalpagina's met kop, 1 uitdraaikop


def test_concept_in_werkmap_en_import_zonder_passende_kolom(tmp_path):
    bron = tmp_path / "werkmap.xlsx"
    bladen = mv.dag_bladen()
    bladen["Concept"] = mv.concept_blad()  # lijstvorm als extra blad in de werkmap
    mv.schrijf(bron, bladen)
    beschikbaarheid(tmp_path / "b.xlsx")
    app = roosterapp.QApplication.instance() or roosterapp.QApplication([])  # noqa: F841
    v = roosterapp.Venster()
    v.open(str(bron))
    assert "Concept 15-12-2026 · 7 lokalen" in v.samenvatting.text()
    gezien = {}

    def kies(dlg):  # 15-12 staat niet in het overzicht: niets voorgevinkt, OK uit
        gezien.update(vink=len(dlg.gekozen()), ok=dlg.ok.isEnabled(), let_op=not dlg.waarschuwing.isHidden(), nodig=dlg.nodig)
        return True
    voor = [list(r) for r in v.bladen["Simu"]]
    assert not v.importeer_simus(str(tmp_path / "b.xlsx"), kies)
    assert gezien == {"vink": 0, "ok": False, "let_op": True, "nodig": 7} and v.bladen["Simu"] == voor
    v.bladen["Simu"] = mv.concept_bladen()["Simu"]
    r = v.maak()
    assert r["fouten"] == [] and v.stat_in.waarde.text() == "61 / 61"
    assert {rij[2] for rij in r["bladen"]["Rooster_Kans1"][1:]} == {"A1", "A2", "B1", "B2", "C1", "Praktijkhuis", "C2"}


def test_wisselen_van_toets_houdt_alles_en_alles_wissen(tmp_path, monkeypatch):
    bron = tmp_path / "werkmap.xlsx"
    mv.schrijf(bron)
    beschikbaarheid(tmp_path / "b.xlsx")
    app = roosterapp.QApplication.instance() or roosterapp.QApplication([])  # noqa: F841
    v = roosterapp.Venster()
    v.open(str(bron))
    v.instellingen.velden["simu_rotatie"].setCurrentText("Na elke toets")
    assert v.importeer_simus(str(tmp_path / "b.xlsx"), lambda dlg: True)
    v.maak()
    v.modus.groep.buttons()[1].click()  # naar Herkansing: meteen een rooster, werkmap en import blijven
    assert v.resultaat and v.resultaat["modus"] == "her" and v.kaart_stapel.currentIndex() == 1
    assert v.bladen is not None and "Simu" in v.ingevoerd
    assert v.instellingen.waarden()["simu_rotatie"] == "Na elke toets"
    v.modus.groep.buttons()[0].click()
    assert v.resultaat["modus"] == "kans1" and v.knop_opslaan.isEnabled()
    monkeypatch.setattr(roosterapp.QMessageBox, "question", lambda *a, **k: roosterapp.QMessageBox.No)
    assert not v.alles_wissen() and v.bladen is not None  # 'Nee': er gebeurt niets
    monkeypatch.setattr(roosterapp.QMessageBox, "question", lambda *a, **k: roosterapp.QMessageBox.Yes)
    assert v.alles_wissen()
    assert v.bladen is None and v.resultaat is None and v.ingevoerd == {} and v.instellingen.waarden() == rc.STANDAARD
    assert v.lbl_bestand.text() == "Geen werkmap geopend" and not v.knop_maak.isEnabled() and not v.knop_wissen.isEnabled()
    assert v.tabel_rooster.rowCount() == 0 and bron.exists()


def test_nieuwe_werkmap_wordt_meteen_geopend(tmp_path):
    app = roosterapp.QApplication.instance() or roosterapp.QApplication([])  # noqa: F841
    v = roosterapp.Venster()
    doel = v.sjabloon(str(tmp_path / "nieuw.xlsx"))
    assert doel and v.pad == doel and v.lbl_bestand.text() == "nieuw.xlsx"


def test_simus_vast_lokaal_in_app(tmp_path):
    bron = tmp_path / "werkmap.xlsx"
    mv.schrijf(bron, mv.concept_bladen())
    app = roosterapp.QApplication.instance() or roosterapp.QApplication([])  # noqa: F841
    v = roosterapp.Venster()
    v.open(str(bron))
    v.instellingen.velden["simu_rotatie"].setCurrentText("Na elke toets")
    assert v.lokalen() == ["A1", "A2", "B1", "B2", "C1", "Praktijkhuis", "C2"]

    def kies(dlg):
        assert [k.currentText() for k in dlg.keuzes][:2] == ["Wisselt mee", "Wisselt mee"]
        assert dlg.keuzes[0].findText("C2") >= 0
        dlg.keuzes[0].setCurrentText("C2")
        dlg.keuzes[1].setCurrentText("Vast (eigen lokaal)")
        dlg.kan[2].setText("vrouw")
        return True
    assert v.simus_instellen(kies)
    r = v.maak()
    rooster = r["bladen"]["Rooster_Kans1"]
    assert {rij[2] for rij in rooster[1:] if rij[6] == "Siem"} == {"C2"} and {rij[2] for rij in rooster[1:] if rij[6] == "Tom"} == {"A1"}
    assert len({rij[6] for rij in rooster[1:] if rij[2] == "B1"}) > 1
    uit = rc.lees_werkmap(v.opslaan(str(tmp_path / "uit.xlsx")))["Simu"]
    assert uit[0][4:6] == ["Vast lokaal", "Kan spelen"] and [r_[4] for r_ in uit[1:3]] == ["C2", "Ja"] and uit[3][5] == "vrouw"
    assert all("vrouw" in rij[3] for rij in rooster[1:] if rij[6] == "Uma")
    # importeren daarna houdt 'Vast lokaal' per naam vast
    nieuw = uitvoer.simu_blad_met(uit, 46371, [{"naam": "Siem", "van": None, "tot": None}, {"naam": "Nieuw", "van": None, "tot": None}])
    assert [r_[4:6] for r_ in nieuw[1:]] == [["C2", ""], ["", ""]]


def test_instellingen_blijven_bewaard_na_afsluiten(tmp_path):
    app = roosterapp.QApplication.instance() or roosterapp.QApplication([])  # noqa: F841
    v = roosterapp.Venster()
    v.instellingen.velden["duur1"].setValue(25)
    v.instellingen.velden["simu_rotatie"].setCurrentText("Na elke toets")
    v.instellingen.velden["pauzes"].setText("12:00-12:30")
    v.close()
    w = roosterapp.Venster()  # programma opnieuw gestart
    assert w.instellingen.waarden()["duur1"] == 25 and w.instellingen.waarden()["simu_rotatie"] == "Na elke toets"
    assert w.instellingen.waarden()["pauzes"] == "12:00-12:30"


def test_werkmap_met_andere_instellingen_overschrijft_eigen_instellingen_niet(tmp_path, monkeypatch):
    bron = tmp_path / "werkmap.xlsx"
    mv.schrijf(bron)  # werkmap met eigen blad Instellingen (standaard: toets 15 min)
    app = roosterapp.QApplication.instance() or roosterapp.QApplication([])  # noqa: F841
    v = roosterapp.Venster()
    v.instellingen.velden["duur1"].setValue(25)
    v.open(str(bron))
    assert v.instellingen.waarden()["duur1"] == 25  # eigen instelling blijft
    assert "Je eigen instellingen zijn actief" in v.statusBar().currentMessage() and v.knop_uit_werkmap.isEnabled()
    assert v.instellingen_uit_werkmap() and v.instellingen.waarden()["duur1"] == 15
    assert roosterapp.Venster().instellingen.waarden()["duur1"] == 15  # en dat wordt ook weer bewaard
    monkeypatch.setattr(roosterapp.QMessageBox, "question", lambda *a, **k: roosterapp.QMessageBox.Yes)
    v.instellingen.velden["duur1"].setValue(30)
    assert v.alles_wissen() and roosterapp.Venster().instellingen.waarden() == rc.STANDAARD


def test_eerste_keer_neemt_instellingen_van_werkmap(tmp_path):
    bron = tmp_path / "werkmap.xlsx"
    b = mv.dag_bladen()
    b["Instellingen"] = rc.instellingen_matrix({"duur1": 22})
    mv.schrijf(bron, b)
    app = roosterapp.QApplication.instance() or roosterapp.QApplication([])  # noqa: F841
    v = roosterapp.Venster()  # nog niets bewaard
    v.open(str(bron))
    assert v.instellingen.waarden()["duur1"] == 22
